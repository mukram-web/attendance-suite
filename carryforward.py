"""
carryforward.py — last week's marks, carried onto this week's roster.

WHY THIS EXISTS
---------------
The pipeline used to rebuild every session column from the pristine roster
export each Monday: 605 attendee reports downloaded, ~429 columns re-marked, so
a fix to the counting rules reached the whole history (the last-10-digit phone
fix repaired three months and +8.4%). The owner has chosen the opposite trade
(2026-09-10): a session is marked ONCE and never recomputed, so a manual upload
of one week's Zoom exports is live in minutes instead of re-doing everything.

Freezing the marks alone is easy — `attendance_core.process_files` only ever
writes the columns it was given files for, so handing it last week's marked
workbook leaves every other column untouched. The problem is that the workbook
also freezes the ROSTER, and every week brings new enrolments, changed payments
and an entirely new batch. So the base workbook for an incremental run is a
MERGE: this week's roster export for structure, last week's workbook for the
session columns.

THE RULES THAT MAKE IT FAITHFUL
-------------------------------
1. **Identity is the marker's identity.** Rows move as students are added and
   removed, so marks are carried per STUDENT, not per row number, matched on
   registered email, then the whole registered number, then its last ten digits
   — `_cell_email` and both branches of `_phone_hit`. A student whose email
   changed but whose phone did not keeps their history, exactly as a re-mark
   would have given it to them. Measured on the live corpus: 61,865 of 61,865
   matched.

2. **A student who joined after a session is 'Absent' for it — but only if they
   were in scope.** A full re-mark writes 'Absent' for any enrolled student who
   is not in that session's attendee list, so filling 'Absent' is what keeps
   every historical number identical the moment the mode changes. (Published
   percentages are `present/denominator` and the denominator is today's
   strength either way, so blank and 'Absent' publish the same numbers; what
   'Absent' additionally protects is `data.build_batch`'s "marked > 30% of
   strength" validity gate, which silently drops a column that falls under it.)
   But a POD session was only offered to that POD — the marker skips everyone
   else and leaves their cell EMPTY — so this fills 'Absent' on a POD column
   only for students whose own POD pref matches it. Ten false absences per
   student per POD day is precisely what that rule prevents.

3. **A column is identified by `attendance_core.session_key`, not by its date.**
   Two sessions on the same day in different years are different sessions; the
   year-blind form exists only for the legacy B17-B28 headers typed by hand.
   Keying everything year-blind collided them, and on a freeze a collided column
   is a session lost for good.

4. **A column this week's export ALSO has is written onto, not duplicated —
   and last week's value wins.** B17-B28 keep their Present/Absent columns
   inside the roster Sheet itself. Leaving those alone looks respectful and is
   wrong: a full run re-marks two of them as a side effect of the shared
   'AI CAP B17 + B21' folders, so last week's workbook holds what the dashboard
   has actually been showing, while the Sheet holds an older understated value.
   Measured: preferring the Sheet moved 82 cells Present -> Absent across two of
   B17's published sessions. Disagreements are counted (`overridden`) so a
   hand-edited Sheet cell shows up in the log instead of vanishing.
"""
from __future__ import annotations

import io

from openpyxl import load_workbook

import attendance_core as ac
import pods

# Session columns start at K on every roster tab, exactly as the marker assumes.
_FIRST_SESSION_COL = 11        # 1-based

_PRESENT = "present"


def _sheet_session_cols(row1, row2, hr) -> dict:
    """{session_key: (col_index_1based, header_value, topic_or_None)}.

    Keyed by `attendance_core.session_key` — the same definition the marker's
    own `dmap` uses — so a column carried here and a column the marker would
    re-mark are recognised as the same session. The header VALUE rides along so
    it can be written out verbatim, preserving the year and the POD that
    `_col_header` put there, and the topic so an `hr == 2` tab keeps its second
    header row.

    Two columns that resolve to the SAME key are a pre-existing anomaly — one
    session with two columns in the workbook. The leftmost wins here, matching
    the marker's own `dmap.setdefault`, so the two agree about which one is the
    session; the duplicate is simply not carried. Under a freeze that means it
    is dropped, which is the right outcome for a column that should not exist,
    but it is not announced. Detecting it properly needs this function to return
    two values and every caller to change, and it has never been observed on the
    live corpus, so it is written down rather than coded around.
    """
    out: dict = {}
    width = max(len(row1 or ()), len(row2 or ()))
    for i in range(_FIRST_SESSION_COL - 1, width):          # 0-based here
        h1 = row1[i] if row1 and i < len(row1) else None
        h2 = row2[i] if (row2 and i < len(row2)) else None
        k = ac.session_key(h1, h2, hr)
        if k:
            out.setdefault(k, (i + 1, h1 if h1 is not None else h2,
                               h2 if hr == 2 else None))
    return out


def _header_row_of(rows) -> int:
    """1-based header row for a tab given its first rows as value tuples."""
    for i, row in enumerate(rows[:3], start=1):
        j = " ".join(str(v or "") for v in (row or ())[:13]).lower()
        if "registered number" in j or "registered mail" in j:
            return i
    return 1


def _col_of(row, *needles) -> int | None:
    """1-based index of the first header cell containing every needle."""
    for i, v in enumerate(row or (), start=1):
        h = str(v or "").strip().lower()
        if h and all(n in h for n in needles):
            return i
    return None


def _identities(email: str, phone: str) -> list:
    """Every key this student can be found under, best first.

    All three of the marker's forms: the email, the whole number, and its last
    ten digits. The full-number branch exists for numbers too short to have a
    last-ten form — drop it and a student with a 9-digit number looks new, and
    their Presents get overwritten with 'Absent' for good.
    """
    out = []
    if email:
        out.append(("e", email))
    if phone:
        out.append(("p", phone))
        if len(phone) >= 10:
            out.append(("p10", phone[-10:]))
    return out


def _merge_marks_into(dst: dict, got: dict) -> None:
    """Fold one row's marks into an identity's marks: PRESENT WINS.

    A person duplicated in the roster with an email on one row and a phone on
    the other legitimately holds two different marks for the same session — the
    marker's rule is an OR over identity, so if either row was Present, they
    were. First-row-wins would make the answer depend on row order.
    """
    for k, v in got.items():
        cur = dst.get(k)
        if cur is None or str(v).strip().lower() == _PRESENT:
            dst[k] = v


def _read_previous(prev_bytes: bytes) -> dict:
    """{batch_key: {"sheet","cols","marks"}} from last week's marked workbook.

    Keyed by `attendance_core._sheet_key`, never by the tab's NAME: the marker
    resolves batches that way precisely because tab names drift ('AI CAP B37'
    becoming 'AI CAP B37 8PM' would otherwise drop that batch's whole history
    without a word).

    Read-only and sparse on purpose: only cells that actually carry a mark are
    kept, so a 60,000-student workbook costs a few MB rather than tens.
    """
    wb = load_workbook(io.BytesIO(prev_bytes), read_only=True, data_only=True)
    out: dict = {}
    try:
        for name in wb.sheetnames:
            key = ac._sheet_key(name)
            if not key or key in out:
                continue
            rows = [tuple(r) for r in wb[name].iter_rows(values_only=True)]
            if not rows:
                continue
            hr = _header_row_of(rows)
            cols = _sheet_session_cols(rows[0], rows[1] if len(rows) > 1 else None, hr)
            if not cols:
                continue
            header = rows[hr - 1]
            rm = _col_of(header, "registered", "mail")
            rn = _col_of(header, "registered", "number")
            marks: dict = {}
            people = 0
            for row in rows[hr:]:
                e = ac._cell_email(row[rm - 1]) if (rm and rm - 1 < len(row)) else ""
                p = ac._cell_phone(row[rn - 1]) if (rn and rn - 1 < len(row)) else ""
                idents = _identities(e, p)
                if not idents:
                    continue
                got = {}
                for k, (ci, _h, _t) in cols.items():
                    v = row[ci - 1] if ci - 1 < len(row) else None
                    s = str(v).strip() if v is not None else ""
                    if s:
                        got[k] = s
                if not got:
                    continue
                # ONE bucket per student, shared by all their identity keys, so
                # "how many students did we fail to place" can be counted by
                # distinct bucket. Keyed per identity instead, a single student
                # counted up to three times.
                bucket = next((marks[i] for i in idents if i in marks), None)
                if bucket is None:
                    bucket = {}
                    people += 1
                for ident in idents:
                    marks.setdefault(ident, bucket)
                _merge_marks_into(bucket, got)
            out[key] = {"sheet": name, "cols": cols, "marks": marks,
                        "people": people}
    finally:
        wb.close()
    return out


def merge_marks(fresh_roster_bytes: bytes,
                prev_marked_bytes: bytes | None) -> tuple[bytes, dict]:
    """This week's roster + last week's session columns -> the base workbook.

    Returns `(xlsx_bytes, report)`. With no previous workbook the fresh export
    is returned untouched and `report["carried"] == 0`, which is exactly a
    normal full run — the escape hatch, and what the very first incremental run
    does.

    The report is printed by the pipeline and kept in the store: a merge that
    silently carried nothing, matched nobody, or quietly dropped a batch is the
    one failure that would otherwise look like a healthy run with a shorter
    history.
    """
    # `_carried_keys` is {batch_key: {session_key}} — every column this merge is
    # responsible for, which is exactly the set the marker must leave alone
    # (`attendance_core.process_files(frozen=...)`). Private, and holding tuple
    # keys that JSON cannot express, so the pipeline pops it before the report
    # goes into the store.
    report = {"tabs": 0, "carried": 0, "cells": 0, "filled": 0, "matched": 0,
              "new_students": 0, "unmatched_prev": 0, "already_present": 0,
              "overridden": 0, "skipped_tabs": [], "warnings": [],
              "_carried_keys": {}}
    if not prev_marked_bytes:
        report["warnings"].append("no previous marked workbook — full run")
        return fresh_roster_bytes, report

    prev = _read_previous(prev_marked_bytes)
    if not prev:
        report["warnings"].append(
            "previous workbook carried no session columns — full run")
        return fresh_roster_bytes, report

    seen_keys = set()
    wb = load_workbook(io.BytesIO(fresh_roster_bytes), data_only=True)
    try:
        for name in wb.sheetnames:
            bkey = ac._sheet_key(name)
            src = prev.get(bkey) if bkey else None
            if not src:
                continue
            seen_keys.add(bkey)
            ws = wb[name]
            hr = ac._header_row(ws)
            rm = ac._col(ws, hr, "registered", "mail")
            rn = ac._col(ws, hr, "registered", "number")
            pc = ac._col(ws, hr, "pod")
            if rm is None and rn is None:
                report["skipped_tabs"].append(name)
                report["warnings"].append(
                    f'Sheet "{name}": no Registered mail/number column — marks '
                    "not carried")
                continue

            # What this week's export already has, keyed like the marker's dmap.
            row1 = tuple(c.value for c in ws[1])
            row2 = tuple(c.value for c in ws[2]) if (ws.max_row or 0) >= 2 else ()
            have = _sheet_session_cols(row1, row2, hr)

            # Ordered by the PREVIOUS workbook's own left-to-right order, so a
            # December column does not sort after January and scramble
            # SessionIdx and the Roster tab across a year boundary.
            todo = sorted(src["cols"].items(), key=lambda kv: kv[1][0])
            # A session this week's export ALSO carries (B17-B28 keep their
            # Present/Absent columns in the roster Sheet itself) is written onto
            # the existing column rather than duplicated. Last week's value wins,
            # and that is deliberate: a full run re-marks some of those legacy
            # columns as a side effect of the shared "AI CAP B17 + B21" folders,
            # so last week's workbook holds what the dashboard has been showing.
            # Leaving the Sheet's own value in place instead dropped 82 Presents
            # off two of B17's published sessions. Where the two disagree it is
            # counted, so a hand-edited Sheet cell shows up in the log rather
            # than disappearing.
            onto = {k: have[k][0] for k, _v in todo if k in have}
            todo = [(k, v) for k, v in todo if k not in have]
            report["already_present"] += len(onto)
            if not todo and not onto:
                continue

            # Merged body cells break a write in openpyxl; the marker unmerges
            # them for the same reason. Header merges (rows 1..hr) stay, and the
            # column maps above were already built from those headers.
            for rng in list(ws.merged_cells.ranges):
                if rng.max_row > hr:
                    ws.unmerge_cells(str(rng))

            nxt = ac._last_used(ws) + 1
            placed = {}
            for key, (_ci, header, topic) in todo:
                ws.cell(1, nxt).value = header
                if hr == 2 and topic:
                    ws.cell(2, nxt).value = topic
                placed[key] = nxt
                nxt += 1
            report["tabs"] += 1
            report["carried"] += len(placed)
            # Both the newly-placed columns and the ones written onto: all of
            # them now hold last week's marks and none may be re-marked.
            report["_carried_keys"][bkey] = set(placed) | set(onto)

            marks = src["marks"]
            used, used_buckets = set(), set()
            for r in range(hr + 1, (ws.max_row or hr) + 1):
                e = ac._cell_email(ws.cell(r, rm).value) if rm else ""
                p = ac._cell_phone(ws.cell(r, rn).value) if rn else ""
                idents = _identities(e, p)
                if not idents:
                    continue
                got = next((marks[i] for i in idents if i in marks), None)
                if got is not None:
                    # Track the BUCKET, not the identity keys. A student whose
                    # email changed is found through their phone, so their old
                    # email key is never touched — counting keys made the merge
                    # warn "their marks are not carried" about precisely the
                    # case the design exists to handle.
                    used_buckets.add(id(got))
                if got is None:
                    report["new_students"] += 1
                    # In scope for a whole-batch column always; for a POD column
                    # only when this is their POD (see the module docstring).
                    rp = None
                    if pc:
                        rp, _multi = pods.from_roster_cell(ws.cell(r, pc).value)
                        rp = rp or pods.UNKNOWN
                    for key, ci in placed.items():
                        if key[1] and rp != key[1]:
                            continue
                        ws.cell(r, ci).value = "Absent"
                        report["filled"] += 1
                    continue
                report["matched"] += 1
                used.update(i for i in idents if i in marks)
                for key, ci in placed.items():
                    v = got.get(key)
                    if v:
                        ws.cell(r, ci).value = v
                        report["cells"] += 1
                # …and onto the columns the export already had. No Absent fill
                # here: those columns already hold a value for everyone this
                # week's export knows about.
                for key, ci in onto.items():
                    v = got.get(key)
                    if not v:
                        continue
                    cur = ws.cell(r, ci).value
                    if str(cur or "").strip() != v:
                        report["overridden"] += 1
                    ws.cell(r, ci).value = v
                    report["cells"] += 1

            # Students whose history could not be placed anywhere on this tab.
            # Counted per tab against that tab's own population - the one number
            # that would reveal "we lost N students' marks", so it must not be a
            # running total compared against a per-tab count. Counted over
            # distinct BUCKETS (one per student) rather than identity keys, so a
            # student found by phone after an email change is not reported lost.
            left = len({id(v) for v in marks.values()} - used_buckets)
            if left:
                report["unmatched_prev"] += left
                report["warnings"].append(
                    f'Sheet "{name}": {left} student(s) in last week\'s workbook '
                    "are not in this week's roster — their marks are not carried")

        missing = [f"{k[0]} B{k[1]}" for k in set(prev) - seen_keys]
        if missing:
            report["warnings"].append(
                "batch tab(s) in last week's workbook but NOT in this week's "
                "roster, so their whole history is dropped: "
                + ", ".join(sorted(missing)))

        out = io.BytesIO()
        wb.save(out)
    finally:
        wb.close()
    return out.getvalue(), report


def summary(report: dict) -> str:
    """One log line the owner can read without opening anything."""
    if not report.get("carried"):
        return "carried nothing forward — " + (
            "; ".join(report.get("warnings") or ["no previous columns"]))
    return (f'{report["carried"]} column(s) carried across {report["tabs"]} tab(s), '
            f'{report["cells"]:,} mark(s) + {report["filled"]:,} filled Absent · '
            f'{report["matched"]:,} student(s) matched, '
            f'{report["new_students"]:,} new, '
            f'{report["unmatched_prev"]:,} not carried · '
            f'{report["already_present"]} column(s) already in this week\'s export')

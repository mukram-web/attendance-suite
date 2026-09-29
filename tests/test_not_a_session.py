"""The not-a-session rule: a webinar L2 lists only as a Hackathon call is never
added, whatever date it sits on.

Owner's ruling, 2026-09-28: the two 26 Sep rooms 89713589616 / 92614513509
"did not run as classes" - both Hackathon calls. Measured 2026-09-30 on the
live L2: 60 webinars carry the word, all of them Intro/Solution calls, none of
them ever a published column. Pinned here:

* the register (`attendance_core.l2_dates`) returns a `NotASession` - a str
  carrying the topic and the dates - for such a webinar;
* a webinar with BOTH a class row and a Hackathon row keeps its class dates;
* the gate (`live_data.l2_gate_reason`) drops it BEFORE the date exemption,
  so a Hackathon room is not marked for the one batch with no class that day;
* a class on the same date is untouched;
* the FFA merge still works beside it.
"""
import datetime
import io
import os
import sys
import unittest

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from openpyxl import Workbook                           # noqa: E402

import attendance_core as ac                            # noqa: E402
import live_data                                        # noqa: E402

DT = datetime.datetime

# A read-only copy of the live L2 workbook, present only on the machine the
# rule was measured on. The regression test skips without it.
REAL_L2 = (r"C:\Users\user\AppData\Local\Temp\claude"
           r"\C--Users-user-OneDrive-Desktop-House-of-Edtech"
           r"\a9bc2c55-796a-419a-9686-44a7734e2ef0\scratchpad\L2.xlsx")


def xlsx(*tabs) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    for title, rows in tabs:
        ws = wb.create_sheet(title)
        for r, row in enumerate(rows, 1):
            for c, v in enumerate(row, 1):
                if v is not None:
                    ws.cell(r, c, v)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def header():
    row = [None] * 17
    row[0], row[1], row[3], row[15] = "Date", "Batch Name", "Topic Name", "Webinar ID"
    return row


def row(date, wid, topic, batch="AI CAP B41"):
    r = [None] * 17
    r[0], r[1], r[3], r[15] = date, batch, topic, wid
    return r


OCTOBER = ("October 2026", [
    header(),
    row(DT(2026, 10, 3), "111 0000 0001", "Prompt Engineering in 2026"),
    row(None, "222 0000 0002", "AI Generalist Hackathon Intro Call", batch="AI CAP B40"),
    row(None, "333 0000 0003", "hackathon solution", batch="AI CAP B39"),
    row(DT(2026, 10, 4), "444 0000 0004", "Hackathon Intro", batch="AI CAP B42"),
    row(None, "444 0000 0004", "Blueprint to Launch", batch="AI CAP B42"),   # same id, a class
    row(None, "555 0000 0005", None, batch="AI CAP B38"),                    # no topic at all
])

L2 = xlsx(OCTOBER)


class TestRegister(unittest.TestCase):
    def setUp(self):
        self.reg = ac.l2_dates(L2, with_ffa=False)

    def test_a_hackathon_webinar_is_registered_as_not_a_session(self):
        v = self.reg["22200000002"]
        self.assertIsInstance(v, ac.NotASession)
        self.assertIsInstance(v, str)
        self.assertEqual(str(v), "AI Generalist Hackathon Intro Call")
        self.assertEqual(v.dates, {"2026_10_03"})

    def test_the_word_alone_is_the_rule_whatever_the_case(self):
        self.assertIsInstance(self.reg["33300000003"], ac.NotASession)

    def test_a_class_row_beats_a_hackathon_row_on_the_same_id(self):
        v = self.reg["44400000004"]
        self.assertIsInstance(v, set)
        self.assertEqual(v, {"2026_10_04"})

    def test_a_class_and_a_topicless_row_are_plain_sets(self):
        self.assertEqual(self.reg["11100000001"], {"2026_10_03"})
        self.assertEqual(self.reg["55500000005"], {"2026_10_04"})

    def test_the_ffa_merge_still_works_beside_it(self):
        reg = ac.l2_dates(L2, with_ffa=True)
        self.assertIsInstance(reg["22200000002"], ac.NotASession)
        self.assertTrue(all(isinstance(v, (set, ac.NotASession)) for v in reg.values()))


class TestGate(unittest.TestCase):
    def setUp(self):
        self.reg = ac.l2_dates(L2, with_ffa=False)

    def test_it_is_dropped_with_the_topic_in_the_reason(self):
        why = live_data.l2_gate_reason("attendee_22200000002_2026_10_03.csv", self.reg)
        self.assertEqual(why, "L2 lists it as 'AI Generalist Hackathon Intro Call', not a session")

    def test_it_is_dropped_even_on_an_exempt_date(self):
        # The date is already marked (a class ran that day): the exemption would
        # let a NEW file through and mark the Hackathon room for B40.
        why = live_data.l2_gate_reason("attendee_22200000002_2026_10_03.csv", self.reg,
                                       exempt_dates={"2026_10_03"})
        self.assertIsNotNone(why)
        self.assertIn("not a session", why)

    def test_a_class_on_the_same_date_still_passes(self):
        self.assertIsNone(live_data.l2_gate_reason("attendee_11100000001_2026_10_03.csv", self.reg))
        self.assertIsNone(live_data.l2_gate_reason("attendee_44400000004_2026_10_04.csv", self.reg))

    def test_the_other_reasons_are_unchanged(self):
        self.assertEqual(live_data.l2_gate_reason("attendee_99900000009_2026_10_03.csv", self.reg),
                         "not in L2")
        self.assertEqual(live_data.l2_gate_reason("attendee_11100000001_2026_10_02.csv", self.reg),
                         "L2 has it on 2026_10_03")

    def test_apply_l2_gate_drops_it_and_keeps_the_class(self):
        items = [("f/attendee_11100000001_2026_10_03.csv", b"x"),
                 ("f/attendee_22200000002_2026_10_03.csv", b"y")]
        kept, dropped = live_data.apply_l2_gate(items, self.reg, set(), stage="fetch")
        self.assertEqual([p for p, _ in kept], ["f/attendee_11100000001_2026_10_03.csv"])
        self.assertEqual(len(dropped), 1)


@unittest.skipUnless(os.path.exists(REAL_L2), "live L2 copy not on this machine")
class TestLiveL2(unittest.TestCase):
    def test_only_hackathon_calls_are_excluded_and_there_are_sixty(self):
        with open(REAL_L2, "rb") as fh:
            reg = ac.l2_dates(fh.read(), with_ffa=False)
        not_sessions = {w: v for w, v in reg.items() if isinstance(v, ac.NotASession)}
        self.assertEqual(len(not_sessions), 60)
        self.assertTrue(all("hackathon" in str(v).lower() for v in not_sessions.values()))
        for wid in ("89713589616", "92614513509"):               # the owner's two
            self.assertIn(wid, not_sessions)
            self.assertEqual(not_sessions[wid].dates, {"2026_09_26"})


if __name__ == "__main__":
    unittest.main()

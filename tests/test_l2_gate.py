"""The L2 gate: a session is added only when L2 registers its webinar ON ITS DATE.

Owner's rule, 2026-09-28: "only add sessions which are in L2 sheet" and "leave
25th sept". The Zoom extracts drive sometimes exports a session into a folder
and filename dated one day early; `_parse_filename` takes the date from the
filename, so the twin mints a phantom Friday column beside the real Saturday
one, and the poll and duration passes - which tie-break on webinar id alone -
let it hijack the real session's rating and duration.

Four things are pinned here: the register (`attendance_core.l2_dates`, which
reads the date column parse_l2 ignores), the predicate
(`live_data.l2_gate_reason`), its NEW-SESSIONS-ONLY scope (a date already
marked in the base workbook is exempt, so history never moves), and the fetch
applying it to the files it returns AND to its failure lists, so a download
that failed on a file the gate drops cannot trip the refuse-to-publish check.
"""
import contextlib
import datetime
import io
import os
import sys
import tempfile
import unittest
from unittest import mock

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from openpyxl import Workbook                           # noqa: E402

import attendance_core as ac                            # noqa: E402
import ffa                                              # noqa: E402
import live_data                                        # noqa: E402

# A read-only copy of the live L2 workbook, present only on the machine the
# regression was measured on. The regression test skips without it.
REAL_L2 = (r"C:\Users\user\AppData\Local\Temp\claude"
           r"\C--Users-user-OneDrive-Desktop-House-of-Edtech"
           r"\a9bc2c55-796a-419a-9686-44a7734e2ef0\scratchpad\L2.xlsx")

DT = datetime.datetime


def xlsx(*tabs) -> bytes:
    """tabs: (title, [row lists]) -> workbook bytes. None cells stay blank."""
    wb = Workbook()
    wb.remove(wb.active)
    for title, rows in tabs:
        ws = wb.create_sheet(title)
        for r, row in enumerate(rows, 1):
            for c, v in enumerate(row, 1):
                if v is not None:
                    ws.cell(r, c, v)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def sep_row(date, wid, batch="AI CAP B35", second_date=None):
    """One row in the September-2026 layout: Date in A, Batch Name in B, Topic
    in D, Webinar ID in P (16), and the unrelated second 'Date' in AF (32)."""
    row = [None] * 33
    row[0], row[1], row[3], row[15], row[31] = date, batch, "A topic", wid, second_date
    return row


def sep_header():
    row = [None] * 33
    row[0], row[1], row[3], row[15], row[31] = ("Date", "Batch Name", "Topic Name",
                                               "Webinar ID", "Date")
    return row


def old_row(date, wid, batch="B19 Alpha Beta"):
    """The Feb-2025 layout: Date A, Batch Name B, Topic C, Webinar ID K (11)."""
    row = [None] * 11
    row[0], row[1], row[2], row[10] = date, batch, "PPT Using AI", wid
    return row


def old_header(first="Date"):
    row = [None] * 11
    row[0], row[1], row[2], row[10] = first, "Batch Name", "Topic Name", "Webinar ID"
    return row


SEPTEMBER = ("September 2026", [
    sep_header(),
    sep_row("1st Week.", None, batch="INDIA/UAE", second_date=DT(2030, 1, 1)),
    sep_row(DT(2026, 9, 26), "914 4100 5879", second_date=DT(2030, 1, 2)),
    sep_row(None, "937 5712 9149"),                    # blank: forward-filled 26th
    sep_row(None, "999 0000 0001", second_date=DT(2030, 1, 3)),
    sep_row("2nd Week.", None, batch="INDIA/UAE"),
    sep_row(DT(2026, 9, 27), "914 4100 5879"),         # the same webinar, second date
    sep_row(None, "888 0000 0002"),
    sep_row("US", None, batch=None),                   # a header that is not a date...
    sep_row(None, "777 0000 0003"),                    # ...so this row has NO date
    sep_row("24th September", "666 0000 0004"),
    sep_row("9/26/2026", "555 0000 0005"),
    sep_row("September 26", "444 0000 0006"),
    sep_row("27 September", "333 0000 0007"),
    sep_row(None, ""),                                 # no webinar id: nothing to register
])

FEB2025 = ("Feb2025", [
    old_header(),
    old_row("23rd Feb, Sunday", "222 0000 0008"),
    old_row("23rd", "111 0000 0009"),                  # a day with no month: no date
])

JAN2026 = ("Jan 2026", [                               # '[' where the header should be
    old_header(first="["),
    old_row("31st Dec", "100 0000 0010"),              # December in a January tab
    old_row("3rd Jan", "100 0000 0011"),
])

DEC2025 = ("Dec 2025", [
    old_header(),
    old_row("1st Jan", "100 0000 0013"),               # January in a December tab
])

NO_YEAR = ("Scratch", [                                # no year anywhere: no date
    old_header(),
    old_row("26th September", "100 0000 0012"),
])

APRIL2026 = ("April 2026", [                           # ' ' where the header should be
    old_header(first=" "),
    old_row("4th April", "100 0000 0014"),
    old_row("Mid Week session", None),
    old_row("10th April", "100 0000 0015"),
])

AUG2026 = ("August 2026", [                            # the ONLY 'Date' header is the stray one
    sep_row(None, "Webinar ID", batch="Batch Name", second_date="Date"),
    sep_row(DT(2026, 8, 1), "100 0000 0016", second_date=DT(2026, 1, 1)),
])

PIVOT = ("Pivot Table 2", [[None], [None]])           # the empty scratch tab that once crashed a build

L2 = xlsx(SEPTEMBER, FEB2025, JAN2026, DEC2025, NO_YEAR, APRIL2026, AUG2026, PIVOT)


class TestL2Dates(unittest.TestCase):
    def setUp(self):
        self.reg = ac.l2_dates(L2, with_ffa=False)

    def test_forward_fill_and_a_webinar_on_two_dates(self):
        self.assertEqual(self.reg["91441005879"], {"2026_09_26", "2026_09_27"})
        self.assertEqual(self.reg["93757129149"], {"2026_09_26"})
        self.assertEqual(self.reg["99900000001"], {"2026_09_26"})
        self.assertEqual(self.reg["88800000002"], {"2026_09_27"})

    def test_a_section_header_ends_the_block_rather_than_inheriting_it(self):
        """'US' is not a date, and the row under it must NOT be dated 27th -
        that is exactly the day-wrong error the gate exists to catch."""
        self.assertEqual(self.reg["77700000003"], set())

    def test_the_second_date_header_is_ignored(self):
        every = {d for v in self.reg.values() for d in v}
        self.assertFalse({d for d in every if d.startswith("2030_")},
                         "column-32 dates leaked into the register")

    def test_the_string_forms(self):
        self.assertEqual(self.reg["66600000004"], {"2026_09_24"})     # 24th September
        self.assertEqual(self.reg["55500000005"], {"2026_09_26"})     # 9/26/2026
        self.assertEqual(self.reg["44400000006"], {"2026_09_26"})     # September 26
        self.assertEqual(self.reg["33300000007"], {"2026_09_27"})     # 27 September
        self.assertEqual(self.reg["22200000008"], {"2025_02_23"})     # 23rd Feb, Sunday

    def test_unparseable_is_no_date_never_a_guess(self):
        self.assertEqual(self.reg["11100000009"], set())              # bare '23rd'
        self.assertEqual(self.reg["10000000012"], set())              # tab has no year

    def test_the_year_is_the_tabs_with_the_dec_jan_wrap(self):
        self.assertEqual(self.reg["10000000010"], {"2025_12_31"})     # '31st Dec' in Jan 2026
        self.assertEqual(self.reg["10000000011"], {"2026_01_03"})
        self.assertEqual(self.reg["10000000013"], {"2026_01_01"})     # '1st Jan' in Dec 2025

    def test_column_a_is_the_date_whatever_its_header_says(self):
        self.assertEqual(self.reg["10000000014"], {"2026_04_04"})     # header ' '
        self.assertEqual(self.reg["10000000015"], {"2026_04_10"})
        self.assertEqual(self.reg["10000000016"], {"2026_08_01"})     # only 'Date' is the stray one

    def test_a_row_with_no_webinar_id_registers_nothing(self):
        self.assertNotIn("", self.reg)

    def test_the_ffa_register_is_folded_in_by_default(self):
        with_ffa = ac.l2_dates(L2)
        for s in ffa.sessions():
            wid, ymd = ffa._norm_wid(s["wid"]), ffa._norm_ymd(s["ymd"])
            self.assertIn(ymd, with_ffa[wid])
            self.assertNotIn(wid, self.reg)

    def test_parse_l2_is_untouched(self):
        """The sibling reads the same tabs; the original must still ignore dates."""
        got = ac.parse_l2(L2)
        self.assertEqual(got["91441005879"], (frozenset({("CAP", 35)}), "A topic"))

    def test_cell_forms_directly(self):
        cases = [
            (DT(2026, 9, 26, 0, 0), None, None, "2026_09_26"),
            (datetime.date(2026, 9, 27), None, None, "2026_09_27"),
            ("2026-09-26 00:00:00", None, None, "2026_09_26"),
            ("9/26/2026", None, None, "2026_09_26"),
            ("26/9/2026", None, None, "2026_09_26"),       # only d/m/y can mean this
            ("26th September", 2026, 9, "2026_09_26"),
            ("2nd  Aug", 2025, 8, "2025_08_02"),
            ("18th April Saturday", 2026, 4, "2026_04_18"),
            ("26th Oct Morning 11 AM", 2025, 10, "2025_10_26"),
            ("18th Oct Mr9", 2025, 10, "2025_10_18"),
            ("27 September", 2026, 9, "2026_09_27"),
            ("September 26", 2026, 9, "2026_09_26"),
            ("Sept 26", 2026, 9, "2026_09_26"),
            ("1st Week.", 2026, 9, None),
            ("4th Week. ", 2026, 9, None),
            ("Last Week.", 2026, 2, None),
            ("US", 2026, 4, None),
            ("Mid Week session", 2026, 4, None),
            ("In House Trainer ", 2026, 6, None),
            ("23rd", 2025, 3, None),
            ("26th September", None, None, None),          # no tab year: no guess
            ("31st Feb", 2026, 2, None),                    # not a real day
            ("L2 Sessions from Aug 2025", 2025, 2, None),
            ("", 2026, 9, None),
            (None, 2026, 9, None),
        ]
        for v, y, m, want in cases:
            with self.subTest(v=v):
                self.assertEqual(ac.l2_cell_date(v, y, m), want)


REG = {"91441005879": {"2026_09_26"}, "93757129149": {"2026_09_26"},
       "77700000003": set()}


class TestPredicate(unittest.TestCase):
    def test_registered_on_its_date_passes(self):
        self.assertIsNone(live_data.l2_gate_reason(
            "attendee_91441005879_2026_09_26.csv", REG))
        self.assertTrue(live_data.l2_registered(
            "2026-09-26 - AI CAP B35 - X/attendee_91441005879_2026_09_26.csv", REG))

    def test_a_day_early_twin_is_dropped_and_says_where_l2_has_it(self):
        self.assertEqual(live_data.l2_gate_reason(
            "2026-09-25 - AI CAP B35 - X/attendee_91441005879_2026_09_25.csv", REG),
            "L2 has it on 2026_09_26")

    def test_a_webinar_l2_never_lists_is_dropped(self):
        self.assertEqual(live_data.l2_gate_reason(
            "attendee_12345678901_2026_09_25.csv", REG), "not in L2")

    def test_a_name_with_no_webinar_id_passes(self):
        for n in ("attendee_Rag beginner.csv", "chat.txt", "Poll Report.csv"):
            with self.subTest(n=n):
                self.assertIsNone(live_data.l2_gate_reason(n, REG))

    def test_a_webinar_l2_lists_without_a_date_passes_on_any_date(self):
        self.assertIsNone(live_data.l2_gate_reason(
            "attendee_77700000003_2026_09_25.csv", REG))

    def test_both_poll_conventions_are_judged_the_same_way(self):
        self.assertEqual(live_data.l2_gate_reason(
            "poll_91441005879_2026_09_25.csv", REG), "L2 has it on 2026_09_26")
        self.assertIsNone(live_data.l2_gate_reason(
            "poll_91441005879_2026_09_26.csv", REG))
        self.assertEqual(live_data.l2_gate_reason(
            "91441005879 - 2026-09-25 - Poll Report.csv", REG), "L2 has it on 2026_09_26")
        self.assertIsNone(live_data.l2_gate_reason(
            "91441005879 - 2026-09-26 - Poll Report.csv", REG))

    def test_no_register_means_no_gate_not_nothing_registered(self):
        for n in ("attendee_12345678901_2026_09_25.csv",
                  "attendee_91441005879_2026_09_25.csv"):
            with self.subTest(n=n):
                self.assertTrue(live_data.l2_registered(n, None))

    def test_apply_logs_every_drop_by_name_and_keeps_the_rest(self):
        items = [("a/attendee_91441005879_2026_09_26.csv", b""),
                 ("b/attendee_91441005879_2026_09_25.csv", b""),
                 ("c/attendee_12345678901_2026_09_25.csv", b""),
                 ("d/attendee_Rag beginner.csv", b"")]
        lines = []
        kept, dropped = live_data.apply_l2_gate(items, REG, stage="fetch",
                                                log=lines.append)
        self.assertEqual([k[0] for k in kept],
                         ["a/attendee_91441005879_2026_09_26.csv",
                          "d/attendee_Rag beginner.csv"])
        self.assertEqual(dropped, [
            ("attendee_91441005879_2026_09_25.csv", "L2 has it on 2026_09_26"),
            ("attendee_12345678901_2026_09_25.csv", "not in L2")])
        self.assertEqual(len(lines), 1)
        self.assertIn("[L2 gate: fetch] 2 file(s) not registered on their date", lines[0])
        self.assertIn("attendee_91441005879_2026_09_25.csv (L2 has it on 2026_09_26)",
                      lines[0])
        self.assertIn("attendee_12345678901_2026_09_25.csv (not in L2)", lines[0])

    def test_apply_reads_listing_dicts_and_bare_names_too(self):
        rows = [{"name": "attendee_91441005879_2026_09_25.csv", "id": "x"},
                {"name": "attendee_91441005879_2026_09_26.csv", "id": "y"}]
        kept, _ = live_data.apply_l2_gate(rows, REG, log=None)
        self.assertEqual([r["id"] for r in kept], ["y"])
        kept, _ = live_data.apply_l2_gate(
            ["attendee_91441005879_2026_09_25.csv", "attendee_x.csv"], REG,
            name_of=lambda n: n, log=None)
        self.assertEqual(kept, ["attendee_x.csv"])

    def test_apply_with_no_register_is_a_no_op_that_prints_nothing(self):
        lines = []
        kept, dropped = live_data.apply_l2_gate(
            [("attendee_12345678901_2026_09_25.csv", b"")], None, log=lines.append)
        self.assertEqual(len(kept), 1)
        self.assertEqual((dropped, lines), ([], []))


# Ten fixed columns, exactly like the real roster, then POD pref, then sessions.
FIXED = ["Contry code", "Registered Number", "Registered mail", "Contry code",
         "Whatsaap Number", "broadcast mail", "batch name", "amount",
         "Payment", "Closing Type", "POD pref"]


def roster(*tabs) -> bytes:
    """tabs: (sheet, [session headers]) -> a one-student roster workbook."""
    return xlsx(*((sheet, [FIXED + list(headers),
                           [91, "919000000001", "a@x.com", 91, "919000000001",
                            "a@x.com", sheet, 1000, "Full Paid", "BDA Closing",
                            "Techies"] + ["Present"] * len(headers)])
                  for sheet, headers in tabs))


class TestNewSessionsOnly(unittest.TestCase):
    """A file dated on a day the base workbook already has a column for is
    exempt - whatever batch, whatever POD - so history can never move."""

    def test_marked_dates_reads_every_tab_pod_and_legacy_form(self):
        got = live_data.marked_dates(roster(
            ("AI CAP B35", ["2026_09_25", "2026_09_26 | Techies"]),
            ("AI CAP B20", ["20th Sep"]),            # a legacy hand-typed header
            ("Attendance", ["2026_01_01"])))         # not a batch tab
        self.assertEqual(got, {"2026_09_25", "2026_09_26", "09_20"})

    def test_a_day_early_file_for_an_already_marked_date_passes(self):
        twin = "attendee_91441005879_2026_09_25.csv"
        self.assertIsNone(live_data.l2_gate_reason(twin, REG, {"2026_09_25"}))
        self.assertIsNone(live_data.l2_gate_reason(twin, REG, {"09_25"}))   # legacy header
        self.assertEqual(live_data.l2_gate_reason(twin, REG, {"2026_09_26"}),
                         "L2 has it on 2026_09_26")

    def test_an_unregistered_webinar_on_a_marked_date_passes_too(self):
        self.assertIsNone(live_data.l2_gate_reason(
            "attendee_12345678901_2026_09_25.csv", REG, {"2026_09_25"}))

    def test_an_empty_exemption_set_exempts_nothing(self):
        self.assertEqual(live_data.l2_gate_reason(
            "attendee_12345678901_2026_09_25.csv", REG, set()), "not in L2")


HEAD = "Attended,First Name,Last Name,Email,Phone\n"


def report(n):
    rows = "".join(f"Yes,A,B,s{i}@x.com,90000000{i:02d}\n" for i in range(n))
    return (HEAD + rows).encode()


F26 = "2026-09-26 - AI CAP B35 - A topic"
F25 = "2026-09-25 - AI CAP B35 - A topic"
GOOD = "attendee_91441005879_2026_09_26.csv"
TWIN = "attendee_91441005879_2026_09_25.csv"
UNKNOWN = "attendee_12345678901_2026_09_25.csv"
TWIN_FAILS = "attendee_93757129149_2026_09_25.csv"


class TestFetchLevel(unittest.TestCase):
    """fetch_new_attendees with Drive faked: two session folders, four files,
    one of which cannot be downloaded."""

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.tree = {
            "ROOT": [{"id": "FA", "name": F26, "mimeType": live_data._FOLDER_MIME},
                     {"id": "FB", "name": F25, "mimeType": live_data._FOLDER_MIME}],
            "FA": [{"id": "f1", "name": GOOD, "mimeType": "text/csv", "size": "1"}],
            "FB": [{"id": "f2", "name": TWIN, "mimeType": "text/csv", "size": "1"},
                   {"id": "f3", "name": UNKNOWN, "mimeType": "text/csv", "size": "1"},
                   {"id": "f4", "name": TWIN_FAILS, "mimeType": "text/csv", "size": "1"}],
        }
        self.bytes = {"f1": report(5), "f2": report(5), "f3": report(2)}
        self.raise_on = {"f4"}
        self.listing_raise_on = set()

        def _list_children(_svc, fid):
            if fid in self.listing_raise_on:
                raise RuntimeError("listing boom")
            return list(self.tree.get(fid, []))

        def _download_any(_svc, fid):
            if fid in self.raise_on:
                raise RuntimeError("download boom")
            return self.bytes[fid]

        for name, val in (("_list_children", _list_children),
                          ("_download_any", _download_any),
                          ("_thread_drive", lambda: None),
                          ("_CACHE_DIR", os.path.join(self.tmp.name, "attendee"))):
            p = mock.patch.object(live_data, name, val)
            p.start()
            self.addCleanup(p.stop)

    def fetch(self, roster_bytes=None, **kw):
        out = io.StringIO()
        with contextlib.redirect_stdout(out):
            files, info = live_data.fetch_new_attendees(
                None, "ROOT", roster_bytes or roster(("AI CAP B35", [])), **kw)
        return files, info, out.getvalue()

    def test_gated_returns_only_the_registered_file_and_clears_its_failures(self):
        files, info, log = self.fetch(l2_dates=REG)
        self.assertEqual([f[0] for f in files], [f"{F26}/{GOOD}"])
        self.assertEqual(info["failed"], 0)
        self.assertEqual(info["download_errors"], [])
        self.assertEqual(info["bad_zips"], [])
        self.assertEqual(info["files"], 1)
        self.assertEqual(info["l2_gate_dropped"], 3)
        self.assertEqual({n for n, _w in info["l2_gate"]}, {TWIN, UNKNOWN, TWIN_FAILS})
        self.assertEqual(info["marked_dates"], set())

    def test_every_drop_is_printed_by_name_with_its_reason(self):
        _files, _info, log = self.fetch(l2_dates=REG)
        self.assertIn("[L2 gate: fetch] 2 file(s) not registered on their date", log)
        self.assertIn(f"{TWIN} (L2 has it on 2026_09_26)", log)
        self.assertIn(f"{UNKNOWN} (not in L2)", log)
        self.assertIn("[L2 gate: fetch, failed downloads] 1 file(s)", log)
        self.assertIn(f"{TWIN_FAILS} (L2 has it on 2026_09_26)", log)

    def test_ungated_is_exactly_the_old_behaviour(self):
        files, info, log = self.fetch()
        self.assertEqual({f[0].rsplit("/", 1)[-1] for f in files}, {GOOD, TWIN, UNKNOWN})
        self.assertEqual(info["failed"], 1)
        self.assertEqual(len(info["download_errors"]), 1)
        self.assertIn(TWIN_FAILS, info["download_errors"][0])
        self.assertEqual((info["l2_gate"], info["l2_gate_dropped"]), ([], 0))
        self.assertNotIn("L2 gate", log)

    def test_a_registered_file_that_fails_to_download_stays_fatal(self):
        self.raise_on = {"f1"}
        _files, info, _log = self.fetch(l2_dates=REG)
        self.assertEqual(info["failed"], 1)
        self.assertIn(GOOD, info["download_errors"][0])

    def test_a_folder_that_cannot_be_listed_stays_fatal(self):
        self.listing_raise_on = {"FB"}
        _files, info, _log = self.fetch(l2_dates=REG)
        self.assertEqual(len(info["listing_errors"]), 1)
        self.assertIn(F25, info["listing_errors"][0])

    def test_a_date_marked_in_any_batch_exempts_the_whole_day(self):
        """B36 already carries a 25th column: the 25th is a loaded weekend, so
        the twin passes and its failure counts - history is never re-judged."""
        base = roster(("AI CAP B35", []), ("AI CAP B36", ["2026_09_25"]))
        files, info, _log = self.fetch(base, l2_dates=REG)
        self.assertEqual({f[0].rsplit("/", 1)[-1] for f in files}, {GOOD, TWIN, UNKNOWN})
        self.assertEqual(info["failed"], 1)
        self.assertEqual(info["marked_dates"], {"2026_09_25"})
        self.assertEqual(info["l2_gate_dropped"], 0)

    def test_an_explicit_exemption_set_overrides_the_workbooks(self):
        files, info, _log = self.fetch(l2_dates=REG, exempt_dates={"2026_09_25"})
        self.assertEqual(len(files), 3)
        self.assertEqual(info["marked_dates"], set())


class FakeSvc:
    """Just enough of the Drive client for the whole-drive listings."""

    def __init__(self, rows):
        self.rows = rows

    def files(self):
        return self

    def list(self, **_kw):
        return self

    def execute(self):
        return {"files": list(self.rows)}


class TestListingGates(unittest.TestCase):
    ROWS = [{"id": "a", "name": GOOD}, {"id": "b", "name": TWIN},
            {"id": "c", "name": UNKNOWN}, {"id": "d", "name": "attendee_Rag beginner.csv"}]
    POLLS = [{"id": "p", "name": "poll_91441005879_2026_09_26.csv"},
             {"id": "q", "name": "poll_91441005879_2026_09_25.csv"},
             {"id": "r", "name": "12345678901 - 2026-09-25 - Poll Report.csv"}]

    def test_list_attendees_and_list_polls_default_to_no_gate(self):
        with mock.patch.object(live_data, "_folder_ids", lambda _f: ["0AXYZ"]):
            self.assertEqual(len(live_data.list_attendees(FakeSvc(self.ROWS), "x")), 4)
            self.assertEqual(len(live_data.list_polls(FakeSvc(self.POLLS), "x")), 3)

    def test_list_attendees_and_list_polls_gate_by_name(self):
        with mock.patch.object(live_data, "_folder_ids", lambda _f: ["0AXYZ"]), \
                contextlib.redirect_stdout(io.StringIO()) as out:
            got = live_data.list_attendees(FakeSvc(self.ROWS), "x", REG)
            self.assertEqual([r["id"] for r in got], ["a", "d"])
            got = live_data.list_polls(FakeSvc(self.POLLS), "x", REG)
            self.assertEqual([r["id"] for r in got], ["p"])
            got = live_data.list_polls(FakeSvc(self.POLLS), "x", REG, {"2026_09_25"})
            self.assertEqual([r["id"] for r in got], ["p", "q", "r"])
        self.assertIn("[L2 gate: attendee listing] 2 file(s)", out.getvalue())
        self.assertIn("[L2 gate: poll listing] 2 file(s)", out.getvalue())

    def test_list_attendee_names_gates_the_names(self):
        rows = [{"name": r["name"]} for r in self.ROWS]
        with mock.patch.object(live_data, "_drive_service", lambda: FakeSvc(rows)), \
                contextlib.redirect_stdout(io.StringIO()):
            self.assertEqual(len(live_data.list_attendee_names("0AXYZ")), 4)
            got = live_data.list_attendee_names("0AXYZ", REG)
            self.assertEqual(got, [GOOD, "attendee_Rag beginner.csv"])
            got = live_data.list_attendee_names("0AXYZ", REG, {"2026_09_25"})
            self.assertEqual(len(got), 4)


@unittest.skipUnless(os.path.exists(REAL_L2), "the live L2 copy is not on this machine")
class TestRealL2Regression(unittest.TestCase):
    """Measured on the live workbook, 2026-09-28: the 26-27 Sep weekend."""

    @classmethod
    def setUpClass(cls):
        with open(REAL_L2, "rb") as fh:
            cls.reg = ac.l2_dates(fh.read(), with_ffa=False)

    def test_the_weekend_registers_22_and_29_webinars(self):
        # 24 and 29 when measured; two of the 26th (89713589616, 92614513509)
        # are Hackathon calls and register as NotASession since 2026-09-30, so
        # they no longer carry a date - see tests/test_not_a_session.py.
        on26 = {w for w, d in self.reg.items() if isinstance(d, set) and "2026_09_26" in d}
        on27 = {w for w, d in self.reg.items() if isinstance(d, set) and "2026_09_27" in d}
        self.assertEqual(len(on26), 22)
        self.assertEqual(len(on27), 29)
        self.assertEqual(len(on26 | on27), 51)
        self.assertIsInstance(self.reg["89713589616"], ac.NotASession)
        self.assertIsInstance(self.reg["92614513509"], ac.NotASession)

    def test_the_two_twinned_webinars_sit_on_saturday_only(self):
        self.assertEqual(self.reg["91441005879"], {"2026_09_26"})
        self.assertEqual(self.reg["93757129149"], {"2026_09_26"})

    def test_the_day_early_twins_are_dropped_and_the_real_exports_kept(self):
        for wid in ("91441005879", "93757129149"):
            with self.subTest(wid=wid):
                self.assertEqual(live_data.l2_gate_reason(
                    f"attendee_{wid}_2026_09_25.csv", self.reg), "L2 has it on 2026_09_26")
                self.assertIsNone(live_data.l2_gate_reason(
                    f"attendee_{wid}_2026_09_26.csv", self.reg))

    def test_almost_every_webinar_in_twenty_tabs_carries_a_date(self):
        """Three rows in Mar 2025 say a bare '23rd'. Nothing else is unparsed."""
        self.assertLessEqual(sum(1 for v in self.reg.values() if not v), 3)
        self.assertGreater(len(self.reg), 1500)


if __name__ == "__main__":
    unittest.main()

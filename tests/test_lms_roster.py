"""Building the roster from the LMS API instead of the Sheet.

The swap is only safe because the workbook it produces is indistinguishable, to
every downstream reader, from Google's export of the Sheet. Four header-row
detectors, three batch-name rules and two column-boundary rules all have to agree
about it, and every one of them fails SILENTLY when it disagrees — a tab the
marker writes to but `data.py` refuses to read produces marks that simply never
reach the dashboard. So most of this file pins the shape, not the logic.

The rest pins the two rules that stop the swap losing data: retained rows survive
(the 9,017 people the API has never heard of), and a blank API closing type never
overwrites a real one.
"""
import io
import os
import re
import sys
import unittest

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from openpyxl import load_workbook, Workbook            # noqa: E402

import attendance_core as ac                            # noqa: E402
import dashboard_core as dc                             # noqa: E402
import data as ddata                                    # noqa: E402
import lms_roster as lr                                 # noqa: E402
import pods                                             # noqa: E402


def customer(email="a@b.com", cc="91", number="9876543210",
             pay="full_paid", close="unknown", alt_emails=None, alt_phones=None):
    return {"unique_learner_id": email, "email": email, "countryCode": cc,
            "number": number, "emailVerified": True,
            "alternativeEmails": alt_emails or [],
            "alternativePhones": alt_phones or [],
            "paymentStatus": pay, "closingType": close, "transactionLedger": []}


def record(name, customers):
    return ({"id": name, "name": name}, customers)


def rows_of(xlsx, sheet):
    wb = load_workbook(io.BytesIO(xlsx))
    ws = wb[sheet]
    return [list(r) for r in ws.iter_rows(values_only=True)]


# ── the column contract ──────────────────────────────────────────────────────

class TestHeaderIsReadableByEveryDetector(unittest.TestCase):
    """Four modules find the header row four different ways. All must succeed."""

    def setUp(self):
        self.xlsx = lr.write_workbook({"B40": [
            lr.customer_row(customer(), "B40", "Finance - AI Career Accelerator Program B40")]})
        self.wb = load_workbook(io.BytesIO(self.xlsx))
        self.ws = self.wb["AI CAP B40"]

    def test_attendance_core_finds_row_1(self):
        self.assertEqual(ac._header_row(self.ws), 1)

    def test_attendance_core_finds_every_column_it_needs(self):
        hr = ac._header_row(self.ws)
        self.assertEqual(ac._col(self.ws, hr, "registered", "mail"), lr.I_MAIL + 1)
        self.assertEqual(ac._col(self.ws, hr, "registered", "number"), lr.I_NUM + 1)
        self.assertEqual(ac._col(self.ws, hr, "whatsa"), lr.I_WA + 1)
        self.assertEqual(ac._col(self.ws, hr, "broadcast"), lr.I_BCAST + 1)
        self.assertEqual(ac._col(self.ws, hr, "pod"), lr.I_POD + 1)

    def test_dashboard_core_exact_equality_detector(self):
        """The weakest of the four: it compares cells for EQUALITY, not substring.

        A header of 'Registered No.' or 'Payment Status' would pass the other
        three detectors and fail this one, silently falling back to row 1 with an
        empty header list.
        """
        tab = dc._Tab(rows=rows_of(self.xlsx, "AI CAP B40"))
        hr, vals = dc._find_header_row(tab)
        self.assertEqual(hr, 1)
        self.assertTrue(vals, "dashboard_core fell back — no header matched")
        self.assertEqual(dc._col_idx(vals, "payment"), lr.I_PAY)
        self.assertEqual(dc._col_idx(vals, "closing type"), lr.I_CLOSE)

    def test_data_layer_detector_and_columns(self):
        rows = rows_of(self.xlsx, "AI CAP B40")
        hr = ddata._find_header_row(rows)
        header = [str(c or "").strip().lower() for c in rows[hr]]
        self.assertEqual(ddata._find_col(header, "registered", "mail"), lr.I_MAIL)
        self.assertEqual(ddata._find_col(header, "payment"), lr.I_PAY)
        self.assertEqual(ddata._find_col(header, "clos"), lr.I_CLOSE)
        self.assertEqual(ddata._find_col(header, "pod"), lr.I_POD)


class TestSessionColumnBoundary(unittest.TestCase):
    def test_pod_column_sits_at_K_and_is_not_a_session(self):
        """POD Prefrence is at K, where session columns begin.

        It survives because `session_key` returns None for it — the same reason
        it is harmless in the live sheet, where B35-B40 all carry it.
        """
        self.assertEqual(lr.I_POD + 1, 11)
        self.assertIsNone(ac.session_key("POD Prefrence", None, 1))

    def test_marker_appends_after_the_pod_column(self):
        xlsx = lr.write_workbook({"B40": [
            lr.customer_row(customer(), "B40", "Finance - AI Career Accelerator Program B40")]})
        ws = load_workbook(io.BytesIO(xlsx))["AI CAP B40"]
        self.assertEqual(ac._last_used(ws), 11)

    def test_pre_pod_batches_omit_the_column_entirely(self):
        """B17-B34 have no POD column in the Sheet, so we must not invent one —
        it would show up as a blank pseudo-session in those batches' grids."""
        xlsx = lr.write_workbook({"B30": [lr.customer_row(customer(), "B30", "")]})
        head = rows_of(xlsx, "AI CAP B30")[0]
        self.assertEqual(len(head), lr.N_FIXED - 1)
        self.assertNotIn("POD Prefrence", head)


class TestTabNames(unittest.TestCase):
    def test_accepted_by_the_strictest_reader(self):
        xlsx = lr.write_workbook({"B7": [lr.customer_row(customer(), "B7", "")]})
        name = load_workbook(io.BytesIO(xlsx)).sheetnames[0]
        self.assertEqual(name, "AI CAP B7")
        self.assertEqual(ddata.batch_label(name), "B7")          # strictest
        self.assertIsNotNone(ac._sheet_key(name))                # loosest
        self.assertEqual(dc._clean_batch_name(name), "B7")
        self.assertRegex(name, r"^\s*AI\s*CAP\s*B\d+\s*$")       # pipeline


# ── API record parsing ───────────────────────────────────────────────────────

class TestBatchNameParsing(unittest.TestCase):
    def test_code_and_pod(self):
        n = "Finance - AI Career Accelerator Program B38"
        self.assertEqual(lr.batch_code(n), "B38")
        self.assertEqual(lr.pod_prefix(n), "Finance")
        self.assertEqual(pods.canon(lr.pod_prefix(n)), "Finance")

    def test_whole_batch_record_has_no_pod(self):
        n = "AI Career Accelerator Program B30"
        self.assertEqual(lr.batch_code(n), "B30")
        self.assertEqual(lr.pod_prefix(n), "")

    def test_other_programmes_are_not_CAP(self):
        """Batch numbers are reused across programme generations — 2025's
        'AICA IC B19' and 2026's CAP B19 share zero of 1,161 people. Matching on
        the number alone would merge two unrelated cohorts."""
        self.assertIsNone(lr.batch_code("AICA IC B19"))
        self.assertIsNone(lr.batch_code("AI Tools Workshop B19"))
        self.assertIsNone(lr.batch_code("Build Side Income Using AI - 6th September"))

    def test_grouping_drops_non_cap(self):
        got = lr.group_batches([
            {"name": "Finance - AI Career Accelerator Program B38"},
            {"name": "Techies - AI Career Accelerator Program B38"},
            {"name": "AI Tools Workshop - 3rd Sept"},
        ])
        self.assertEqual(sorted(got), ["B38"])
        self.assertEqual(len(got["B38"]), 2)


class TestValueMapping(unittest.TestCase):
    def test_payment_maps_to_the_rosters_own_spelling(self):
        for api, want, active in [("full_paid", "Full Paid", True),
                                  ("booking", "Booking Amount", True),
                                  ("partially_paid", "Partially Paid", True),
                                  ("refunded", "Unidentified/Refunded", False),
                                  ("none", "", False)]:
            row = lr.customer_row(customer(pay=api), "B40", "")
            self.assertEqual(row[lr.I_PAY], want)
            # All three independent active-vocabularies must agree with us.
            self.assertEqual(ddata.is_active(row[lr.I_PAY]), active, api)
            self.assertEqual(dc._is_active(row[lr.I_PAY]), active, api)

    def test_l3_purchased_no_longer_leaks_through_raw(self):
        row = lr.customer_row(customer(close="l3_purchased"), "B40", "")
        self.assertEqual(row[lr.I_CLOSE], "L3 Purchased")

    def test_unknown_closing_is_blank_not_a_category(self):
        row = lr.customer_row(customer(close="unknown"), "B40", "")
        self.assertEqual(row[lr.I_CLOSE], "")

    def test_phone_is_text_and_carries_the_country_code(self):
        row = lr.customer_row(customer(cc="91", number="9876543210"), "B40", "")
        self.assertIsInstance(row[lr.I_NUM], str)
        self.assertEqual(row[lr.I_NUM], "919876543210")
        # and the marker's last-10 rule still finds the bare Zoom form
        self.assertTrue(ac._phone_hit(ac._cell_phone(row[lr.I_NUM]),
                                      set(), {"9876543210"}))

    def test_broadcast_mail_excludes_the_primary(self):
        row = lr.customer_row(
            customer(email="a@b.com", alt_emails=["a@b.com", "c@d.com"]), "B40", "")
        self.assertEqual(row[lr.I_BCAST], "c@d.com")


class TestApiRows(unittest.TestCase):
    def test_one_person_in_two_pods_appears_once(self):
        c = customer(email="dup@x.com")
        rows, warns = lr.api_rows("B38", [
            record("Finance - AI Career Accelerator Program B38", [c]),
            record("Techies - AI Career Accelerator Program B38", [c]),
        ])
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0][lr.I_POD],
                         "Finance - AI Career Accelerator Program B38")
        self.assertEqual(warns, [])

    def test_unrecognised_pod_warns_rather_than_guessing(self):
        rows, warns = lr.api_rows("B38", [
            record("Astrologers - AI Career Accelerator Program B38", [customer()])])
        self.assertEqual(len(rows), 1)
        self.assertTrue(any("Astrologers" in w for w in warns), warns)

    def test_the_real_pod_spellings_all_canon(self):
        """Every POD prefix the API actually returns must resolve, including the
        'Enterpreneurs' typo, or students land in Unassigned."""
        for p in ["Finance", "Techies", "Educators", "Students", "Healthcare",
                  "Data", "Generalist", "Content Creators", "Sales/Marketing/HR",
                  "Business Owners/Enterpreneurs", "Ops/Supply Chain"]:
            self.assertIsNotNone(
                pods.canon(p), f"POD {p!r} does not canon")


# ── the three layers ─────────────────────────────────────────────────────────

def prev_workbook(tabs):
    """A stand-in for last week's marked workbook: fixed columns + session cols."""
    wb = Workbook()
    wb.remove(wb.active)
    for sheet, students in tabs.items():
        ws = wb.create_sheet(sheet)
        ws.append(lr.HEADERS + ["2026_08_23"])
        for mail, phone, pay, close, pod in students:
            ws.append(["91", phone, mail, "", "", "", sheet, "1000",
                       pay, close, pod, "Present"])
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


class TestReadPrevious(unittest.TestCase):
    def test_reads_attributes_and_ignores_session_columns(self):
        prev = prev_workbook({"AI CAP B38": [
            ("a@b.com", "919876543210", "Full Paid", "BDA Closing", "Finance")]})
        got = lr.read_previous(prev)
        self.assertEqual(list(got), ["B38"])
        row = got["B38"][0]
        self.assertEqual(len(row), lr.N_FIXED)          # no session column leaked
        self.assertEqual(row[lr.I_MAIL], "a@b.com")
        self.assertEqual(row[lr.I_CLOSE], "BDA Closing")
        self.assertEqual(row[lr.I_AMT], "1000")

    def test_skips_non_roster_tabs(self):
        prev = prev_workbook({"AI CAP B38": [("a@b.com", "91987", "Full Paid", "", "")],
                              "Attendance summary": [("x@y.com", "91", "", "", "")],
                              "Pivot Table 2": []})
        self.assertEqual(list(lr.read_previous(prev)), ["B38"])

    def test_no_previous_workbook_is_not_an_error(self):
        self.assertEqual(lr.read_previous(b""), {})
        self.assertEqual(lr.read_previous(None), {})

    def test_retains_a_loosely_named_tab_the_marker_would_still_mark(self):
        """`AI CAP B37 8PM` is a real shape. carryforward reads it with
        ac._sheet_key and would carry its marks onto our clean `AI CAP B37`,
        so failing to retain its students would make every one of them
        `unmatched_prev` — marks binned, and GATE 6 structurally blind to it."""
        prev = prev_workbook({"AI CAP B37 8PM": [
            ("late@x.com", "919876543210", "Full Paid", "System", "")]})
        got = lr.read_previous(prev)
        self.assertEqual(list(got), ["B37"])
        self.assertEqual(got["B37"][0][lr.I_MAIL], "late@x.com")

    def test_two_tabs_for_one_batch_merge_and_warn(self):
        prev = prev_workbook({
            "AI CAP B37": [("a@x.com", "919000000001", "Full Paid", "System", "")],
            "AI CAP B37 8PM": [("a@x.com", "919000000001", "Full Paid", "System", ""),
                               ("b@x.com", "919000000002", "Full Paid", "System", "")],
        })
        warns = []
        got = lr.read_previous(prev, warns)
        self.assertEqual(sorted(r[lr.I_MAIL] for r in got["B37"]),
                         ["a@x.com", "b@x.com"])       # duplicate collapsed
        self.assertTrue(any("B37" in w for w in warns), warns)

    def test_bsiai_tabs_are_not_cap(self):
        """BSIAI has its own roster and its own batch numbers; reading a
        `BSIAI B1` tab as CAP B1 would collide two programmes head-on."""
        prev = prev_workbook({"BSIAI B1": [("x@y.com", "91900", "Full Paid", "", "")]})
        self.assertEqual(lr.read_previous(prev), {})

    def test_a_scratch_tab_is_not_read_as_a_batch(self):
        """`_sheet_key` finds a digit in 'Sheet1' and `_track` DEFAULTS to CAP,
        so a stray tab with roster-shaped columns would otherwise arrive as a
        batch B1 that does not exist and show up on the dashboard."""
        prev = prev_workbook({"Sheet1": [("a@b.com", "919876543210", "Full Paid", "", "")],
                              "AI CAP B38": [("c@d.com", "919876543211", "Full Paid", "", "")]})
        self.assertEqual(list(lr.read_previous(prev)), ["B38"])

    def test_empty_scratch_tabs_are_ignored(self):
        """People add pivot/scratch tabs to the roster; one already crashed a
        build. 'Pivot Table 2' even looks like batch 2 to the loose rule, so the
        no-identity-columns check is what actually rejects it."""
        prev = prev_workbook({"AI CAP B38": [("a@b.com", "91987", "Full Paid", "", "")],
                              "Pivot Table 2": []})
        self.assertEqual(list(lr.read_previous(prev)), ["B38"])


class TestMerge(unittest.TestCase):
    def test_person_the_api_never_heard_of_survives(self):
        """The 9,017. Losing them is the whole risk of this migration."""
        retained = lr.read_previous(prev_workbook({"AI CAP B38": [
            ("ghost@x.com", "919999999999", "Full Paid", "BDA Closing", "Finance")]}))["B38"]
        rows, st = lr.merge(retained, [lr.customer_row(customer(email="new@x.com"), "B38", "")])
        mails = [r[lr.I_MAIL] for r in rows]
        self.assertIn("ghost@x.com", mails)
        self.assertIn("new@x.com", mails)
        self.assertEqual(st["added"], 1)
        self.assertEqual(st["retained"], 1)

    def test_blank_api_closing_never_overwrites_a_real_one(self):
        """The measured regression: the API collapses BDA Closing and Old
        Customer into `unknown`, which maps to blank. B40 reads 58% Unknown
        because of it. A retained value must win."""
        retained = lr.read_previous(prev_workbook({"AI CAP B38": [
            ("a@b.com", "919876543210", "Full Paid", "BDA Closing", "Finance")]}))["B38"]
        fresh = [lr.customer_row(customer(email="a@b.com", close="unknown",
                                          pay="refunded"), "B38", "")]
        rows, st = lr.merge(retained, fresh)
        self.assertEqual(len(rows), 1)
        self.assertEqual(st["updated"], 1)
        self.assertEqual(rows[0][lr.I_CLOSE], "BDA Closing")   # preserved
        self.assertEqual(rows[0][lr.I_PAY], "Unidentified/Refunded")  # refreshed

    def test_real_api_closing_does_overwrite(self):
        retained = lr.read_previous(prev_workbook({"AI CAP B38": [
            ("a@b.com", "919876543210", "Full Paid", "", "Finance")]}))["B38"]
        fresh = [lr.customer_row(customer(email="a@b.com", close="bda_collection"), "B38", "")]
        rows, _ = lr.merge(retained, fresh)
        self.assertEqual(rows[0][lr.I_CLOSE], "BDA Collection")

    def test_matches_on_phone_when_the_email_changed(self):
        retained = lr.read_previous(prev_workbook({"AI CAP B38": [
            ("old@x.com", "919876543210", "Full Paid", "System", "Finance")]}))["B38"]
        fresh = [lr.customer_row(
            customer(email="new@x.com", cc="91", number="9876543210"), "B38", "")]
        rows, st = lr.merge(retained, fresh)
        self.assertEqual(len(rows), 1, "last-10 phone match failed — row duplicated")
        self.assertEqual(st["updated"], 1)
        self.assertEqual(rows[0][lr.I_MAIL], "old@x.com")   # identity untouched

    def test_retained_identity_is_never_rewritten(self):
        """carryforward keys marks to the retained email/phone. Changing either
        here would orphan that person's history."""
        retained = lr.read_previous(prev_workbook({"AI CAP B38": [
            ("a@b.com", "919876543210", "Full Paid", "System", "Finance")]}))["B38"]
        before = list(retained[0])
        fresh = [lr.customer_row(customer(email="a@b.com", number="1111111111"), "B38", "")]
        rows, _ = lr.merge(retained, fresh)
        self.assertEqual(rows[0][lr.I_NUM], before[lr.I_NUM])
        self.assertEqual(rows[0][lr.I_MAIL], before[lr.I_MAIL])

    def test_amount_is_preserved_because_the_api_has_none(self):
        retained = lr.read_previous(prev_workbook({"AI CAP B38": [
            ("a@b.com", "919876543210", "Full Paid", "System", "Finance")]}))["B38"]
        rows, _ = lr.merge(retained, [lr.customer_row(customer(email="a@b.com"), "B38", "")])
        self.assertEqual(rows[0][lr.I_AMT], "1000")


class TestBuildLayering(unittest.TestCase):
    def setUp(self):
        self.prev = prev_workbook({
            "AI CAP B38": [("old38@x.com", "919000000038", "Full Paid", "BDA Closing", "Finance")],
            "AI CAP B39": [("old39@x.com", "919000000039", "Full Paid", "Old Customer", "Techies")],
        })
        self.api = {
            "B38": [record("Finance - AI Career Accelerator Program B38",
                           [customer(email="api38@x.com", number="9000000138")])],
            "B39": [record("Techies - AI Career Accelerator Program B39",
                           [customer(email="api39@x.com", number="9000000139")])],
            "B40": [record("Finance - AI Career Accelerator Program B40",
                           [customer(email="api40@x.com", number="9000000140")])],
        }

    def test_frozen_batch_takes_no_api_rows(self):
        xlsx, rep = lr.build(self.prev, self.api, live=["B39", "B40"])
        mails = [r[lr.I_MAIL] for r in rows_of(xlsx, "AI CAP B38")[1:]]
        self.assertEqual(mails, ["old38@x.com"])
        self.assertIn("B38", rep["frozen"])

    def test_live_batch_is_refreshed_and_keeps_its_retained_people(self):
        xlsx, rep = lr.build(self.prev, self.api, live=["B39", "B40"])
        mails = sorted(r[lr.I_MAIL] for r in rows_of(xlsx, "AI CAP B39")[1:])
        self.assertEqual(mails, ["api39@x.com", "old39@x.com"])
        self.assertIn("B39", rep["refreshed"])

    def test_brand_new_batch_is_seeded_even_though_it_is_not_live(self):
        xlsx, rep = lr.build(self.prev, self.api, live=["B39"])
        self.assertIn("AI CAP B40", load_workbook(io.BytesIO(xlsx)).sheetnames)
        self.assertIn("B40", rep["seeded"])

    def test_tabs_are_ordered_numerically(self):
        xlsx, _ = lr.build(self.prev, self.api, live=["B39", "B40"])
        self.assertEqual(load_workbook(io.BytesIO(xlsx)).sheetnames,
                         ["AI CAP B38", "AI CAP B39", "AI CAP B40"])

    def test_no_previous_workbook_seeds_everything(self):
        xlsx, rep = lr.build(None, self.api, live=["B40"])
        self.assertEqual(sorted(rep["seeded"]), ["B38", "B39", "B40"])


class TestRefreshHonesty(unittest.TestCase):
    def test_a_batch_wanted_live_but_with_no_api_rows_counts_as_frozen(self):
        prev = prev_workbook({"AI CAP B41": [
            ("a@x.com", "919000000041", "Full Paid", "System", "Finance")]})
        retained = lr.read_previous(prev)
        # live names B41, but api_by_code has nothing for it
        _x, rep = lr.build_from(retained, {}, live=["B41"])
        self.assertIn("B41", rep["frozen"])
        self.assertNotIn("B41", rep["refreshed"])
        self.assertTrue(any("B41" in w for w in rep["warnings"]), rep["warnings"])


class TestFetchPlanning(unittest.TestCase):
    def test_live_is_the_two_highest_numbers(self):
        self.assertEqual(lr.live_codes(["B7", "B38", "B40", "B39"]), ["B39", "B40"])

    def test_sorts_numerically_not_lexically(self):
        self.assertEqual(lr.live_codes(["B9", "B10", "B8"], n=2), ["B9", "B10"])

    def test_plan_covers_live_plus_unseen_and_nothing_else(self):
        got = lr.plan_fetch(api_codes=["B38", "B39", "B40", "B41"],
                            retained_codes=["B38", "B39", "B40"],
                            live=["B39", "B40"])
        self.assertEqual(got, ["B39", "B40", "B41"])

    def test_the_lms_own_back_catalogue_is_not_mistaken_for_new(self):
        """The LMS carries CAP batches back to B1; the roster has only ever
        tracked B17+. Without a floor, 'anything we have not seen' makes sixteen
        ancient cohorts look brand new and seeds them onto the dashboard as empty
        batches — and costs sixteen slow API fetches to do it."""
        got = lr.plan_fetch(api_codes=[f"B{i}" for i in range(1, 42)],
                            retained_codes=[f"B{i}" for i in range(17, 41)],
                            live=["B40"])
        self.assertEqual(got, ["B40", "B41"])

    def test_bootstrap_with_nothing_retained_starts_at_B17(self):
        got = lr.plan_fetch(api_codes=[f"B{i}" for i in range(1, 42)],
                            retained_codes=[], live=["B40", "B41"])
        self.assertEqual(got[0], "B17")
        self.assertNotIn("B16", got)
        self.assertEqual(len(got), 25)              # B17..B41

    def test_a_live_batch_the_api_lacks_is_reported_not_swallowed(self):
        """The current cohort is the one whose enrolment is still moving. If the
        API has no record for it, plan_fetch used to drop it silently and
        build_from still called it "refreshed" — so it sat on last week's
        roster while every log said otherwise."""
        missed = []
        got = lr.plan_fetch(api_codes=["B39", "B40"],
                            retained_codes=["B39", "B40", "B41"],
                            live=["B40", "B41"], missing=missed)
        self.assertEqual(got, ["B40"])
        self.assertEqual(missed, ["B41"])

    def test_frozen_history_costs_no_api_calls(self):
        """The reason the refresh policy exists: a full sweep is ~95 batch
        records and the endpoint 504s under load."""
        got = lr.plan_fetch(api_codes=[f"B{i}" for i in range(17, 42)],
                            retained_codes=[f"B{i}" for i in range(17, 42)],
                            live=["B40", "B41"])
        self.assertEqual(got, ["B40", "B41"])


if __name__ == "__main__":
    unittest.main()

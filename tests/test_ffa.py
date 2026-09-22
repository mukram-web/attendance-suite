"""FFA: marked only from the exports the owner hands over, whole batch only.

FFA is the one session type L2 does not register — 0 of 104 rows carry a
Webinar ID — so `ffa.py` is its register instead, and the owner's two rules
(2026-09-23) are what these tests hold in place:

  1. only a (webinar, date) that was handed over is marked, which matters
     because FFA B20's four days share ONE webinar id and only two of them are
     AI CAP slots; and
  2. FFA marks the WHOLE batch, never a domain, because it is one room for
     everyone.

Rule 2 has teeth: a POD column's denominator is a fraction of the batch, so
filing a room the whole batch attended under a domain would publish several
hundred percent.
"""
import io
import os
import sys
import unittest

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import attendance_core as ac  # noqa: E402
import ffa  # noqa: E402
import pods  # noqa: E402
import sheets as dsheets  # noqa: E402

WID = "91643507137"            # FFA B20, days 1-4 — one id, four dates
INVITE_WID = "96664517207"     # the 2 Sep invitation call — a different id
MARKED = "2026_09_12"          # day 3: an AI CAP slot, handed over
UNMARKED = "2026_09_11"        # day 2: same webinar, a Friday, NOT handed over

# B33 and B39 students. B39 is in the room too (3-4% of it did attend FFA), and
# must NOT be marked: it is not one of the batches the owner registered.
B33 = ("b33@x.com", "919000000033")
B39 = ("b39@x.com", "919000000039")


def report_csv(rows, wid=WID, date="09/12/2026"):
    """A minimal Zoom Attendee Report. `rows` is [(email, phone)]."""
    out = ["Attendee Report", "",
           "Topic,Webinar ID,Actual Start Time,Unique Viewers",
           f"4-Day Financial Freedom Accelerator,{wid},{date} 18:43,2", "",
           "Attendee Details",
           "Attended,First Name,Last Name,Email,Phone,Join Time,Leave Time,"
           "Time in Session (minutes)"]
    for em, ph in rows:
        out.append(f"Yes,A,B,{em},{ph},{date} 18:45,{date} 22:50,247")
    return "\n".join(out).encode()


class TestRegistry(unittest.TestCase):
    """What ffa.lookup will and will not claim."""

    def test_a_handed_over_date_names_the_six_batches(self):
        keys, topic, label = ffa.lookup(WID, MARKED)
        self.assertEqual(keys, frozenset(("CAP", n) for n in (33, 34, 35, 36, 37, 38)))
        self.assertEqual(topic, "FFA")
        self.assertIn("B33", label)

    def test_another_date_of_the_SAME_webinar_is_not_claimed(self):
        # The whole point of keying on the date: days 1-2 ran on this same id.
        # Dropping their exports on the drive must change nothing.
        self.assertIsNone(ffa.lookup(WID, UNMARKED))
        self.assertIsNone(ffa.lookup(WID, "2026_09_10"))

    def test_the_invite_call_is_not_claimed(self):
        self.assertIsNone(ffa.lookup(INVITE_WID, "2026_09_02"))

    def test_an_ordinary_webinar_is_not_claimed(self):
        self.assertIsNone(ffa.lookup("91695866411", "2026_09_19"))

    def test_zooms_spacing_of_a_webinar_id_still_matches(self):
        # The report header writes '916 4350 7137'; filenames write the digits.
        self.assertIsNotNone(ffa.lookup("916 4350 7137", "2026-09-12"))


class TestLabelIsWholeBatch(unittest.TestCase):
    """The label FFA presents itself with must read as a whole-batch session."""

    def test_the_label_names_no_domain(self):
        _k, _t, label = ffa.lookup(WID, MARKED)
        self.assertEqual(pods.from_l2_label(label), pods.WHOLE_BATCH,
                         "a ' - <domain>' tail here would file FFA under a POD")

    def test_every_registered_batch_is_readable_from_the_label(self):
        keys, _t, label = ffa.lookup(WID, MARKED)
        self.assertEqual(set(ac.extract_batches(label)), set(keys))


class MarkerBase(unittest.TestCase):
    def _roster(self):
        from openpyxl import Workbook
        wb = Workbook(); wb.remove(wb.active)
        for tab, (em, ph), b in (("AI CAP B33", B33, "B33"), ("AI CAP B39", B39, "B39")):
            ws = wb.create_sheet(tab)
            ws.append(["Country", "Registered Number", "Registered Mail", "WhatsApp",
                       "Broadcast", "Batch", "Amount", "Payment", "Close Type",
                       "POD Prefrence"])
            ws.append([91, ph, em, "", "", b, 0, "Full Paid", "BDA Closing", "Techies"])
        buf = io.BytesIO(); wb.save(buf); return buf.getvalue()

    def _l2(self):
        """A real schedule that does NOT list the FFA webinar — the normal case.

        Non-empty on purpose: `process_files` only enforces "L2 is the register"
        when a schedule actually loaded, so an empty one would let the FFA
        export through by the folder-name fallback and prove nothing.
        """
        from openpyxl import Workbook
        wb = Workbook(); wb.remove(wb.active)
        ws = wb.create_sheet("Sep 2026")
        ws.append(["Date", "Webinar ID", "Batch Name", "Topic"])
        ws.append(["09/19/2026", "91695866411", "AI CAP B33", "Some Session"])
        buf = io.BytesIO(); wb.save(buf); return buf.getvalue()

    def _run(self, files):
        _, report, warns = ac.process_files(self._roster(), self._l2(), files,
                                            values_only=True)
        return report, warns


class TestMarking(MarkerBase):
    def test_a_handed_over_ffa_export_is_marked_although_L2_never_lists_it(self):
        report, _ = self._run([
            (f"2026-09-12 - AI CAP B33 , B34 - FFA/attendee_{WID}_2026_09_12.csv",
             report_csv([B33, B39]))])
        rows = {r["batch"]: r for r in report}
        self.assertIn("AI CAP B33", rows, "the registered batch is marked")
        self.assertEqual(rows["AI CAP B33"]["present"], 1)
        self.assertEqual(rows["AI CAP B33"]["topic"], "FFA")

    def test_it_is_a_WHOLE_BATCH_column(self):
        report, _ = self._run([
            (f"2026-09-12 - AI CAP B33 - FFA/attendee_{WID}_2026_09_12.csv",
             report_csv([B33]))])
        self.assertEqual([r["pod"] for r in report], [""])

    def test_a_folder_naming_a_domain_cannot_make_it_a_POD_column(self):
        # The safety net for rule 2. `pods.from_folder` reads this name and
        # would otherwise hand FFA a Techies-sized denominator.
        report, _ = self._run([
            (f"2026-09-12 - AI CAP B33 - Techies - FFA/attendee_{WID}_2026_09_12.csv",
             report_csv([B33]))])
        self.assertEqual([r["pod"] for r in report], [""])

    def test_a_batch_that_was_not_registered_is_left_alone(self):
        # B39 sat in the FFA room (3-4% of it really did) but is not one of the
        # six batches handed over, so it gets no column at all.
        report, _ = self._run([
            (f"2026-09-12 - FFA/attendee_{WID}_2026_09_12.csv",
             report_csv([B33, B39]))])
        self.assertEqual({r["batch"] for r in report}, {"AI CAP B33"})

    def test_an_unregistered_date_of_the_same_webinar_marks_nothing(self):
        # Rule 1, end to end: day 2's export on the drive changes nothing.
        report, warns = self._run([
            (f"2026-09-11 - FFA/attendee_{WID}_2026_09_11.csv",
             report_csv([B33], date="09/11/2026"))])
        self.assertEqual(report, [])
        self.assertTrue(any("not in the L2 schedule" in w for w in warns),
                        "and it is reported as unregistered, not silently dropped")


class TestDashboardSeesIt(unittest.TestCase):
    """`data.REQUIRE_L2` hides any column with no topic, so FFA needs one."""

    def _l2(self):
        from openpyxl import Workbook
        wb = Workbook(); wb.remove(wb.active)
        ws = wb.create_sheet("Sep 2026")
        ws.append(["Date", "Webinar ID", "Batch Name", "Topic"])
        ws.append(["09/19/2026", "91695866411", "AI CAP B33", "Some Session"])
        buf = io.BytesIO(); wb.save(buf); return buf.getvalue()

    def test_the_lookup_gives_every_registered_batch_a_topic(self):
        lookup, labels = dsheets.webinar_topic_lookup(
            [(f"2026-09-12 - FFA/attendee_{WID}_2026_09_12.csv", b"")],
            self._l2(), with_labels=True)
        for b in ("B33", "B34", "B35", "B36", "B37", "B38"):
            self.assertEqual(lookup.get((b, "09_12", "")), "FFA", b)
            self.assertIn("B33", labels[(b, "09_12", "")])

    def test_it_is_keyed_whole_batch_so_a_POD_column_cannot_borrow_it(self):
        lookup = dsheets.webinar_topic_lookup(
            [(f"2026-09-12 - FFA/attendee_{WID}_2026_09_12.csv", b"")], self._l2())
        self.assertNotIn(("B33", "09_12", "Techies"), lookup)

    def test_an_unregistered_date_gets_no_topic(self):
        lookup = dsheets.webinar_topic_lookup(
            [(f"2026-09-11 - FFA/attendee_{WID}_2026_09_11.csv", b"")], self._l2())
        self.assertEqual(lookup, {})


if __name__ == "__main__":
    unittest.main()

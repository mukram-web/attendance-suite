"""Domain-POD tests: name reconciliation, per-POD denominators, date rollup."""
import os
import sys
import unittest

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import data  # noqa: E402
import pods  # noqa: E402
from tests.test_data import make_tab  # noqa: E402


class TestCanon(unittest.TestCase):
    def test_every_roster_spelling_seen_in_b35_b38(self):
        for raw, want in [
            ("AI Generalist", "Generalist"),
            ("Techies", "Techies"),
            ("Finance", "Finance"),
            ("Business Owners/Enterpreneurs", "Business Owners"),
            ("Sales/Marketing/HR", "Sales/Marketing/HR"),
            ("Operations/Supply Chain", "Ops/Supply Chain"),
            ("Educators", "Educators"),
            ("Students", "Students"),
            ("Healthcare", "Healthcare"),
            ("Data", "Data"),
            ("Content Creators", "Content Creators"),
            ("Techies - AI Career Accelerator Program B35", "Techies"),
        ]:
            self.assertEqual(pods.from_roster_cell(raw)[0], want, raw)

    def test_blank_and_unidentified_are_a_real_bucket(self):
        for raw in ("", None, "Unidentified", "   "):
            self.assertEqual(pods.from_roster_cell(raw)[0], pods.UNKNOWN, repr(raw))

    def test_multi_valued_cell_takes_the_last_and_flags_itself(self):
        got, multi = pods.from_roster_cell(
            "Finance - AI Career Accelerator Program B36; Business Owners/Enterpreneurs")
        self.assertEqual(got, "Business Owners")
        self.assertTrue(multi)

    def test_every_l2_spelling_seen_on_b35_plus(self):
        for raw, want in [
            ("AI CAP B35 - Techies", "Techies"),
            ("AI CAP B35 - Generalist", "Generalist"),
            ("AICAPB35, B36-S/M/HR", "Sales/Marketing/HR"),
            ("AICAPB35, B36-Ops/SC", "Ops/Supply Chain"),
            ("AICAPB35, B36-BusinessOwners", "Business Owners"),
            ("AICAPB35, B36-ContentCreators", "Content Creators"),
            ("AI CAP B35 - Ops/Supply Chain", "Ops/Supply Chain"),
        ]:
            self.assertEqual(pods.from_l2_label(raw), want, raw)

    def test_a_label_naming_no_domain_is_a_whole_batch_session(self):
        for raw in ("AI CAP B35", "AI CAP B17", "B37-42", "AI CAP B35 - All Domains",
                    "AICAPB37- Common"):
            self.assertEqual(pods.from_l2_label(raw), pods.WHOLE_BATCH, raw)

    def test_general_is_the_whole_batch_but_generalist_is_the_pod(self):
        """B36 uses both, so they cannot mean the same thing: 'General' on 22 Aug
        (whole batch) and 'Generalist' on 30 Aug (the AI Generalist POD, one of
        eleven that day). Folding them together scored a whole-batch session
        against a POD - B37's 22 Aug read 1,631 present out of 571."""
        self.assertEqual(pods.from_l2_label("AI CAP B36 - General"), pods.WHOLE_BATCH)
        self.assertEqual(pods.from_l2_label("AI CAP B37 - General"), pods.WHOLE_BATCH)
        self.assertEqual(pods.from_l2_label("AICAPB35, B36-Generalist"), "Generalist")
        self.assertEqual(pods.from_roster_cell("AI Generalist")[0], "Generalist")

    def test_an_unknown_domain_is_none_so_the_caller_can_warn(self):
        """Never invent a bucket: a new POD silently folded into another would
        move a denominator with nothing on screen to say so."""
        self.assertIsNone(pods.from_l2_label("AI CAP B35 - Robotics"))
        self.assertIsNone(pods.canon("Robotics"))

    def test_a_topic_in_the_domain_slot_is_not_read_as_a_domain(self):
        self.assertIsNone(pods.from_folder(
            "2026-08-02 - AI CAP B33 - Office Productivity & Communication"))
        self.assertEqual(pods.from_folder(
            "2026-08-23 - AI CAP B35 - Techies - Python with AI"), "Techies")


class TestRunTogetherFolderNames(unittest.TestCase):
    """'AICAPB37- Techies' has no space before the dash and no word boundary
    before B37. The folder resolved to NO batch, so the session was never even
    downloaded and L2 was never consulted - it vanished with no warning."""

    def test_run_together_folder_resolves_to_its_batch(self):
        import attendance_core as ac
        for name, want in [
            ("2026-08-30 - AICAPB37- Techies - Cursor AI: Build Software", ("CAP", 37)),
            ("2026-08-30 - AICAPB38- AllDomains - Your AI Employee at work", ("CAP", 38)),
            ("2026-08-30 - AICAPB37- Common - Excel with AI", ("CAP", 37)),
        ]:
            self.assertIn(want, ac._folder_batches(name), name)

    def test_spaced_forms_still_resolve(self):
        import attendance_core as ac
        self.assertEqual(ac._folder_batches(
            "2026-08-23 - AI CAP B35 - Techies - Python with AI"), {("CAP", 35)})
        self.assertEqual(ac._folder_batches(
            "2026-08-02 - AI CAP B33 - Office Productivity"), {("CAP", 33)})


def _pod_tab(pod_rows, marks_by_col, headers):
    """A roster tab with a POD column and one column per (date, POD) session.

    pod_rows      : list of POD cell values, one per student
    marks_by_col  : list of columns, each a list of 'Present'/'Absent'/'' per student
    headers       : row-1 header for each session column
    """
    head = ["CountryCode", "Registered Number", "Registered mail", "CountryCode",
            "Whatsaap Number", "broadcast mail", "batch name", "Amount",
            "Payment", "Closing Type", "POD pref"] + headers
    rows = [head]
    for i, pod in enumerate(pod_rows):
        r = [91, 900000000 + i, f"s{i}@x.com", 91, 900000000 + i, f"s{i}@x.com",
             "AI CAP B35", 1000, "Full Paid", "BDA Closing", pod]
        r += [col[i] for col in marks_by_col]
        rows.append(r)
    return rows


class TestPerPodDenominator(unittest.TestCase):
    """10 students: 6 Techies, 4 Generalist. One Techies-only session on 23 Aug
    that 3 of the 6 attended."""

    def setUp(self):
        self.pod_rows = ["Techies"] * 6 + ["AI Generalist"] * 4
        marks = [["Present"] * 3 + ["Absent"] * 3 + [""] * 4]
        self.rows = _pod_tab(self.pod_rows, marks, ["2026_08_23 | Techies"])

    def test_a_pod_session_is_scored_against_that_pod_not_the_batch(self):
        b = data.build_batch(self.rows, "B35", {("B35", "08_23"): "Python with AI"})
        s = b["sessions"][0]
        self.assertEqual(s["pod"], "Techies")
        self.assertEqual(s["total"], 6)        # the POD, not the 10-strong batch
        self.assertEqual(s["present"], 3)
        self.assertEqual(s["pct"], 50.0)       # 3/6, NOT 3/10 = 30%
        self.assertEqual(s["absent"], 3)

    def test_pod_membership_is_reported_per_batch(self):
        b = data.build_batch(self.rows, "B35", {("B35", "08_23"): "T"})
        self.assertEqual(b["pods"]["Techies"]["strength"], 6)
        self.assertEqual(b["pods"]["Generalist"]["strength"], 4)

    def test_a_small_pod_session_survives_the_validity_threshold(self):
        """Thresholds scale to the POD. Against the batch, a 2-member POD could
        never clear '30% of strength marked' and its session would vanish."""
        pod_rows = ["Data"] * 2 + ["AI Generalist"] * 18
        marks = [["Present", "Absent"] + [""] * 18]
        rows = _pod_tab(pod_rows, marks, ["2026_08_23 | Data"])
        b = data.build_batch(rows, "B35", {("B35", "08_23"): "VBA with AI"})
        self.assertEqual(b["n_sessions"], 1)
        self.assertEqual(b["sessions"][0]["total"], 2)
        self.assertEqual(b["sessions"][0]["pct"], 50.0)


class TestDateRollup(unittest.TestCase):
    """Two PODs meeting the same day: the batch line must weight by POD size."""

    def setUp(self):
        pod_rows = ["Techies"] * 6 + ["AI Generalist"] * 4
        techies = ["Present"] * 3 + ["Absent"] * 3 + [""] * 4      # 3 of 6
        general = [""] * 6 + ["Present"] * 4                       # 4 of 4
        self.rows = _pod_tab(pod_rows, [techies, general],
                             ["2026_08_23 | Techies", "2026_08_23 | Generalist"])
        self.l2 = {("B35", "08_23"): "Same-day sessions"}

    def test_both_sessions_survive_instead_of_merging_into_one_column(self):
        b = data.build_batch(self.rows, "B35", self.l2)
        self.assertEqual(b["n_sessions"], 2)
        self.assertEqual(sorted(s["pod"] for s in b["sessions"]),
                         ["Generalist", "Techies"])

    def test_the_batch_headline_is_weighted_by_pod_size(self):
        b = data.build_batch(self.rows, "B35", self.l2)
        self.assertEqual(len(b["by_date"]), 1)
        d = b["by_date"][0]
        self.assertEqual(d["present"], 7)       # 3 Techies + 4 Generalist
        self.assertEqual(d["n_pods"], 2)
        self.assertEqual(d["pct"], 70.0)        # 7 of the 10-strong batch
        self.assertEqual(b["avg_pct"], 70.0)    # NOT the 75% mean of 50% and 100%


class TestPodPercentCannotExceedItsPod(unittest.TestCase):
    def test_marks_outside_the_pod_do_not_inflate_the_numerator(self):
        """B37's 22 Aug read 1,631 present out of 571 - 285.6% - because the
        column carried marks for the whole batch while the denominator was the
        POD. Numerator and denominator must span the same people."""
        pod_rows = ["AI Generalist"] * 4 + ["Techies"] * 6
        # a Generalist-labelled column that (wrongly) marks everyone present
        marks = [["Present"] * 10]
        rows = _pod_tab(pod_rows, marks, ["2026_08_22 | Generalist"])
        b = data.build_batch(rows, "B37", {("B37", "08_22"): "Prompt Engineering"})
        s = b["sessions"][0]
        self.assertEqual(s["total"], 4)
        self.assertEqual(s["present"], 4)      # the 4 Generalists, not all 10
        self.assertEqual(s["pct"], 100.0)
        self.assertLessEqual(s["pct"], 100.0)

    def test_no_session_anywhere_can_report_over_100_percent(self):
        pod_rows = ["Data"] * 2 + ["AI Generalist"] * 18
        rows = _pod_tab(pod_rows, [["Present"] * 20], ["2026_08_23 | Data"])
        b = data.build_batch(rows, "B35", {("B35", "08_23"): "x"})
        for s in b["sessions"]:
            self.assertLessEqual(s["present"], s["total"], s)
            self.assertLessEqual(s["pct"], 100.0, s)


class TestOnlyInvitedStudentsCount(unittest.TestCase):
    def test_a_day_only_one_pod_met_is_scored_against_that_pod(self):
        """B35's 22 Aug: the Generalist POD alone met. The other 2,291 students
        were never invited, so they are not absentees - dividing by the whole
        batch reported 12% for a session 41% of its POD attended."""
        pod_rows = ["AI Generalist"] * 4 + ["Techies"] * 6
        marks = [["Present", "Present", "Absent", "Absent"] + [""] * 6]
        rows = _pod_tab(pod_rows, marks, ["2026_08_22 | Generalist"])
        b = data.build_batch(rows, "B35", {("B35", "08_22"): "Agents Part 1"})
        d = b["by_date"][0]
        self.assertEqual(d["total"], 4)     # the Generalist POD, not all 10
        self.assertEqual(d["present"], 2)
        self.assertEqual(d["pct"], 50.0)    # not 2/10 = 20%

    def test_someone_in_both_a_pod_and_a_whole_batch_session_counts_once(self):
        """A Techies student attending both that day's whole-batch session and
        their POD session is one person present, not two."""
        pod_rows = ["Techies"] * 4 + ["AI Generalist"] * 6
        whole = ["Present"] * 4 + ["Absent"] * 6          # everyone invited
        techies = ["Present"] * 4 + [""] * 6              # same 4 again
        rows = _pod_tab(pod_rows, [whole, techies],
                        ["2026_08_15", "2026_08_15 | Techies"])
        b = data.build_batch(rows, "B35", {("B35", "08_15"): "x"})
        d = b["by_date"][0]
        self.assertEqual(d["total"], 10)    # a whole-batch session invites all
        self.assertEqual(d["present"], 4)   # NOT 8
        self.assertEqual(d["pct"], 40.0)


class TestClosingTypeIsPerPerson(unittest.TestCase):
    def test_a_student_is_scored_only_on_sessions_their_pod_ran(self):
        """6 Techies, 4 Generalist, one session each. Every student attended
        their own. Dividing by all 2 sessions would report 50%; the honest
        number is 100% - nobody could attend the other POD's session."""
        pod_rows = ["Techies"] * 6 + ["AI Generalist"] * 4
        techies = ["Present"] * 6 + [""] * 4
        general = [""] * 6 + ["Present"] * 4
        rows = _pod_tab(pod_rows, [techies, general],
                        ["2026_08_23 | Techies", "2026_08_23 | Generalist"])
        b = data.build_batch(rows, "B35", {("B35", "08_23"): "x"})
        self.assertEqual(len(b["closing"]), 1)            # all 'BDA Closing'
        self.assertEqual(b["closing"][0]["att"], 100.0)


class TestPodViewHandlesProgrammesWithoutPods(unittest.TestCase):
    """A programme with no `by_date` rollup and no PODs. pod_view assumed both
    and returned an empty session list, blanking the whole tab with
    "No sessions logged for None yet"."""

    def test_a_batch_with_no_by_date_is_returned_untouched(self):
        import dash_view
        d = {"code": "B1", "strength": 10, "active": 10, "n_sessions": 1,
             "avg_pct": 50.0, "peak": 50.0, "low": 50.0, "closing": [],
             "sessions": [{"col": None, "mm": "08_15", "date_lbl": "15 Aug",
                           "topic": "x", "present": 5, "absent": 5, "total": 10,
                           "pct": 50.0, "present_only": False, "no_l2": False}]}
        self.assertEqual(dash_view.pod_names(d), [])
        v = dash_view.pod_view(d, None)
        self.assertEqual(len(v["sessions"]), 1)
        self.assertEqual(v["sessions"][0]["date_lbl"], "15 Aug")


class TestPodViewRowContract(unittest.TestCase):
    """Every row pod_view emits is rendered by the same table as a real session,
    so it must carry the same keys. A rollup row missing "absent" crashed the
    whole Dashboard tab with a KeyError."""

    KEYS = {"date_lbl", "topic", "present", "absent", "total", "pct",
            "present_only", "no_l2"}

    def test_all_pods_rollup_rows_carry_every_key_the_table_reads(self):
        import dash_view
        pod_rows = ["Techies"] * 6 + ["AI Generalist"] * 4
        techies = ["Present"] * 3 + ["Absent"] * 3 + [""] * 4
        general = [""] * 6 + ["Present"] * 4
        rows = _pod_tab(pod_rows, [techies, general],
                        ["2026_08_23 | Techies", "2026_08_23 | Generalist"])
        d = data.build_batch(rows, "B35", {("B35", "08_23"): "x"})
        for sel in [None] + dash_view.pod_names(d):
            for row in dash_view.pod_view(d, sel)["sessions"]:
                self.assertFalse(self.KEYS - set(row),
                                 f"pod={sel!r} missing {sorted(self.KEYS - set(row))}")
            dash_view.sessions_table_html(dash_view.pod_view(d, sel))

    def test_rollup_absent_is_consistent_with_present_and_total(self):
        import dash_view
        pod_rows = ["Techies"] * 6 + ["AI Generalist"] * 4
        rows = _pod_tab(pod_rows, [["Present"] * 3 + ["Absent"] * 3 + [""] * 4],
                        ["2026_08_23 | Techies"])
        d = data.build_batch(rows, "B35", {("B35", "08_23"): "x"})
        row = dash_view.pod_view(d, None)["sessions"][0]
        self.assertEqual(row["present"] + row["absent"], row["total"])


class TestNoPodBatchIsUnchanged(unittest.TestCase):
    def test_a_batch_with_no_pod_column_behaves_exactly_as_before(self):
        """The regression guarantee for B17-B34: no POD column, no PODs, and the
        weighted rollup collapses to one session per date - the old number."""
        rows = make_tab(
            [{"mail": f"s{i}@x.com", "payment": "Full Paid", "closing": "BDA Closing",
              "marks": ["Present" if i < 6 else "Absent"]} for i in range(10)],
            session_dates=["2026_04_04"], session_topics=["Topic"],
        )
        b = data.build_batch(rows, "B17", {("B17", "04_04"): "Topic"})
        self.assertEqual(b["pods"], {})
        self.assertEqual(b["sessions"][0]["pod"], "")
        self.assertEqual(b["sessions"][0]["total"], 10)
        self.assertEqual(b["avg_pct"], 60.0)


if __name__ == "__main__":
    unittest.main()


class TestPollRatings(unittest.TestCase):
    """Ratings come from the session's own feedback poll, joined on Webinar ID."""

    WIDE = ("Overview\n"
            "Generate Time,Meeting Topic,Meeting/Webinar ID,Actual Start Time\n"
            "08/29/2026,Demand Forecasting,92569378808,08/29/2026\n\n"
            "Launched Polls\n#,Poll Name,Questions,Responses\n1,Feedback,4,3\n\n"
            "Feedback\n"
            "#,User Name,Email Address,Submitted Date and Time,"
            "What was your overall session feedback? ,How would you rate the trainer?,"
            "How likely would you recommend it to your friends?,"
            "What can we improve in future sessions? (Description)\n"
            "1,A,a@x.com,08/29/2026,5,5,5,all good\n"
            "2,B,b@x.com,08/29/2026,4,3,4,more practice\n"
            "3,C,c@x.com,08/29/2026,3,4,2,\n")

    def test_wide_export_averages_each_question(self):
        import polls
        r = polls.parse(self.WIDE)
        self.assertEqual(r["responses"], 3)
        self.assertAlmostEqual(r["session"], 4.0)
        self.assertAlmostEqual(r["trainer"], 4.0)
        self.assertAlmostEqual(r["recommend"], 3.67, places=2)

    def test_free_text_column_is_never_read_as_a_score(self):
        import polls
        self.assertIsNone(polls._classify(
            "What can we improve in future sessions? (Description)"))

    def test_question_wording_variants_all_classify(self):
        import polls
        for q, want in [
            ("What was your overall session feedback?", "session"),
            ("How was your overall experience in today's project session?", "session"),
            ("How satisfied are you with today's project session?", "session"),
            ("How would you rate the trainer? Please note: 1 = Very Poor", "trainer"),
            ("How likely would you recommend it to your friends?", "recommend"),
        ]:
            self.assertEqual(polls._classify(q), want, q)

    def test_out_of_range_values_are_dropped(self):
        import polls
        self.assertIsNone(polls._score("great"))
        self.assertIsNone(polls._score("9"))
        self.assertEqual(polls._score("5 - Excellent"), 5.0)

    def test_duplicate_exports_keep_the_one_with_more_responses(self):
        import polls
        small = self.WIDE
        big = self.WIDE + "4,D,d@x.com,08/29/2026,5,5,5,\n"
        got = polls.parse_files([("poll_111_2026_08_29.csv", small.encode()),
                                 ("poll_111_2026_08_29.csv", big.encode())])
        self.assertEqual(got["111"]["responses"], 4)

    def test_a_poll_with_no_rating_question_yields_nothing(self):
        import polls
        txt = ("Overview\n\nLaunched Polls\n#,Poll Name,Questions,Responses\n"
               "1,Quiz,1,2\n\nQuiz\n#,User Name,Email Address,"
               "Submitted Date and Time,Which tool did you use?\n"
               "1,A,a@x.com,08/29/2026,Excel\n")
        self.assertEqual(polls.parse_files([("poll_222_2026_08_29.csv", txt.encode())]), {})


class TestRatingSurvivesTheRollup(unittest.TestCase):
    """The default view rolls sessions up per date. The rating has to travel with
    them - it was dropped there, so every AI CAP session showed a dash while the
    data was sitting in the store."""

    def _batch(self):
        pod_rows = ["Techies"] * 6 + ["AI Generalist"] * 4
        techies = ["Present"] * 3 + ["Absent"] * 3 + [""] * 4
        general = [""] * 6 + ["Present"] * 4
        rows = _pod_tab(pod_rows, [techies, general],
                        ["2026_08_23 | Techies", "2026_08_23 | Generalist"])
        ratings = {("B35", "08_23", "Techies"): {"session": 4.0, "trainer": 4.1,
                                                 "recommend": 4.2, "responses": 100},
                   ("B35", "08_23", "Generalist"): {"session": 5.0, "trainer": 5.0,
                                                    "recommend": 5.0, "responses": 300}}
        return data.build_batch(rows, "B35", {("B35", "08_23"): "x"}, None, ratings)

    def test_ratings_reach_the_individual_sessions(self):
        b = self._batch()
        got = {s["pod"]: s["rating"] for s in b["sessions"]}
        self.assertEqual(got, {"Techies": 4.0, "Generalist": 5.0})

    def test_rollup_row_averages_by_responses_not_by_session(self):
        import dash_view
        row = dash_view.pod_view(self._batch(), None)["sessions"][0]
        # (4.0*100 + 5.0*300) / 400 = 4.75, NOT the flat mean of 4.5
        self.assertAlmostEqual(row["rating"], 4.75, places=2)
        self.assertEqual(row["rating_n"], 400)

    def test_a_date_with_no_poll_shows_no_rating_rather_than_zero(self):
        import dash_view
        pod_rows = ["Techies"] * 6 + ["AI Generalist"] * 4
        rows = _pod_tab(pod_rows, [["Present"] * 3 + ["Absent"] * 3 + [""] * 4],
                        ["2026_08_23 | Techies"])
        b = data.build_batch(rows, "B35", {("B35", "08_23"): "x"})
        for sel in [None, "Techies"]:
            for s in dash_view.pod_view(b, sel)["sessions"]:
                self.assertIsNone(s["rating"])
                self.assertIn("rating", s)


class TestPollReportExportShape(unittest.TestCase):
    """Zoom's "Poll Report" export puts one ROW per answer behind a header that
    LOOKS wide ('#, User Name, User Email, Submitted Time, Question, Answer').
    It defeated both other readers, so B38's 30 Aug showed no rating while a
    657-response poll sat on Drive."""

    REPORT = chr(10).join([
        "Poll Report",
        "Report generated time,08/30/2026 06:18:23 PM",
        "",
        "Topic,Webinar ID,Actual Start Time,# Poll",
        "Your AI Employee at work,99299344465,08/30/2026 11:00:00 AM,1190",
        "",
        "#,User Name,User Email,Submitted Time (America/Los_Angeles),Question,Answer",
        "1,A,a@x.com,08/30/2026,have you open claude ?,Yes",
        "2,A,a@x.com,08/30/2026,What was your overall session feedback? ,4",
        "3,A,a@x.com,08/30/2026,How would you rate the trainer?,5",
        "4,A,a@x.com,08/30/2026,What can we improve in future sessions? (Description),NA",
        "5,B,b@x.com,08/30/2026,What was your overall session feedback? ,2",
        "6,B,b@x.com,08/30/2026,How would you rate the trainer?,3",
    ])

    def test_ratings_are_read_from_the_row_per_answer_shape(self):
        import polls
        r = polls.parse(self.REPORT)
        self.assertAlmostEqual(r["session"], 3.0)     # (4 + 2) / 2
        self.assertAlmostEqual(r["trainer"], 4.0)     # (5 + 3) / 2

    def test_respondents_are_counted_as_people_not_answers(self):
        import polls
        self.assertEqual(polls.parse(self.REPORT)["responses"], 2)

    def test_non_rating_questions_in_that_shape_are_ignored(self):
        import polls
        self.assertIsNone(polls.parse(self.REPORT)["recommend"])

    def test_alternate_rating_wording_is_recognised(self):
        import polls
        self.assertEqual(polls._classify("Please rate the today's class"), "session")


class TestPollFileNaming(unittest.TestCase):
    """Two conventions exist on Drive. Reading only Zoom's own left 59
    webinars' feedback stranded."""

    def test_both_conventions_yield_the_same_key(self):
        import polls
        a = polls.name_key("poll_99299344465_2026_08_30.csv")
        b = polls.name_key("99299344465 - 2026-08-30 - Poll Report-2 (1).csv")
        self.assertEqual(a, ("99299344465", "08_30"))
        self.assertEqual(a, b)

    def test_a_name_with_no_webinar_id_is_ignored(self):
        import polls
        self.assertIsNone(polls.name_key("report_poll (2).csv"))
        self.assertIsNone(polls.name_key("Session_Links.txt"))

    def test_a_folder_path_prefix_does_not_break_the_key(self):
        import polls
        self.assertEqual(
            polls.name_key("2026-08-30 - AI CAP B38/poll_99299344465_2026_08_30.csv"),
            ("99299344465", "08_30"))


class TestNoPollLabel(unittest.TestCase):
    def test_an_unrated_session_says_so_rather_than_showing_a_dash(self):
        import dash_view
        html = dash_view._rating({"rating": None})
        self.assertIn("no poll conducted", html)

    def test_a_rated_session_shows_the_score_and_response_count(self):
        import dash_view
        html = dash_view._rating({"rating": 4.56, "rating_n": 657,
                                  "rating_trainer": 4.75, "rating_recommend": 4.51})
        self.assertIn("4.6", html)
        self.assertIn("657", html)
        self.assertNotIn("no poll conducted", html)


class TestComplementRoom(unittest.TestCase):
    """A batch's early weekends run TWO PARALLEL ROOMS: one domain POD, and one
    for everybody else. That second room has no name anywhere - not in the
    roster, not in L2 - so it arrives as an unlabelled column and used to be
    scored against full batch strength. B40's 12 Sep reported 1,611 attendees as
    43% of 3,711 when they were 51% of the 3,150 people actually invited.

    The shape is IDENTICAL to a genuine whole-batch session that a POD also met
    alongside, so the two can only be told apart by behaviour: if the POD met
    instead of the main room, its members are Absent there and Present in their
    own.
    """

    def test_two_parallel_rooms_score_the_complement_not_the_whole_batch(self):
        pod_rows = ["Techies"] * 4 + ["AI Generalist"] * 6
        # the Techies sat out the main room entirely - they were in their own
        common = ["Absent"] * 4 + ["Present"] * 3 + ["Absent"] * 3
        techies = ["Present"] * 3 + ["Absent"] + [""] * 6
        rows = _pod_tab(pod_rows, [common, techies],
                        ["2026_08_15", "2026_08_15 | Techies"])
        b = data.build_batch(rows, "B35", {})
        by_pod = {s["pod"]: s for s in b["sessions"] if s.get("mm")}
        self.assertEqual(by_pod["Common"]["total"], 6)   # NOT 10
        self.assertEqual(by_pod["Common"]["present"], 3)
        self.assertEqual(by_pod["Common"]["pct"], 50.0)  # not 3/10 = 30%
        self.assertEqual(by_pod["Techies"]["total"], 4)
        self.assertEqual(by_pod["Techies"]["present"], 3)

    def test_a_genuine_whole_batch_session_is_left_alone(self):
        """The regression guard. When the POD's members are Present in the main
        room too, the batch really did all meet - that column must keep full
        strength as its denominator."""
        pod_rows = ["Techies"] * 4 + ["AI Generalist"] * 6
        whole = ["Present"] * 4 + ["Absent"] * 6
        techies = ["Present"] * 4 + [""] * 6
        rows = _pod_tab(pod_rows, [whole, techies],
                        ["2026_08_15", "2026_08_15 | Techies"])
        b = data.build_batch(rows, "B35", {})
        whole_sx = [s for s in b["sessions"] if s.get("mm") and not s.get("excl")]
        self.assertTrue(any(s["total"] == 10 for s in whole_sx), whole_sx)

    def test_a_few_people_in_the_wrong_room_do_not_flip_it(self):
        """One or two students wander in. That is noise, not a signal, and it
        must not turn a real whole-batch session into a complement one."""
        pod_rows = ["Techies"] * 10 + ["AI Generalist"] * 10
        whole = ["Present"] * 8 + ["Absent"] * 2 + ["Present"] * 10
        techies = ["Present"] * 9 + ["Absent"] + [""] * 10
        rows = _pod_tab(pod_rows, [whole, techies],
                        ["2026_08_15", "2026_08_15 | Techies"])
        b = data.build_batch(rows, "B35", {})
        self.assertTrue(any(s["total"] == 20 for s in b["sessions"] if s.get("mm")))

    def test_the_complement_carries_which_pods_it_excluded(self):
        pod_rows = ["Techies"] * 4 + ["AI Generalist"] * 6
        common = ["Absent"] * 4 + ["Present"] * 6
        techies = ["Present"] * 4 + [""] * 6
        rows = _pod_tab(pod_rows, [common, techies],
                        ["2026_08_15", "2026_08_15 | Techies"])
        b = data.build_batch(rows, "B35", {})
        comp = [s for s in b["sessions"] if s.get("pod") == "Common"]
        self.assertEqual(len(comp), 1)
        self.assertEqual(comp[0]["excl"], ["Techies"])


class TestBatchPayloadIsSerialisable(unittest.TestCase):
    """`build_store` JSON-encodes this payload straight into the DuckDB store.
    A stray set survives every unit test that only reads the numbers and then
    kills the weekly run at [6/8] - which is exactly how a `set` in the per-date
    rollup reached production once."""

    def test_a_two_room_weekend_json_encodes(self):
        import json
        pod_rows = ["Techies"] * 4 + ["AI Generalist"] * 6
        common = ["Absent"] * 4 + ["Present"] * 6
        techies = ["Present"] * 4 + [""] * 6
        rows = _pod_tab(pod_rows, [common, techies],
                        ["2026_08_15", "2026_08_15 | Techies"])
        b = data.build_batch(rows, "B35", {})
        json.dumps(b)          # must not raise


class TestCommonIsSelectable(unittest.TestCase):
    """The complement room is a property of the SESSION, not of any student, so
    it never appears in d["pods"] and the POD filter could not offer it. For a
    two-room batch that made the split invisible: 'All PODs' blended the two,
    and picking 'Generalist' found no sessions at all."""

    def _two_room_batch(self):
        pod_rows = ["Techies"] * 4 + ["AI Generalist"] * 6
        common = ["Absent"] * 4 + ["Present"] * 3 + ["Absent"] * 3
        techies = ["Present"] * 3 + ["Absent"] + [""] * 6
        rows = _pod_tab(pod_rows, [common, techies],
                        ["2026_08_15", "2026_08_15 | Techies"])
        return data.build_batch(rows, "B35", {})

    def test_pod_view_returns_only_the_complement_sessions(self):
        import dash_view
        import pods as _p
        v = dash_view.pod_view(self._two_room_batch(), _p.COMMON)
        self.assertEqual(len(v["sessions"]), 1)
        self.assertEqual(v["sessions"][0]["pod"], _p.COMMON)
        self.assertEqual(v["sessions"][0]["total"], 6)
        self.assertEqual(v["strength"], 6)      # not the batch's 10
        self.assertEqual(v["avg_pct"], 50.0)

    def test_the_techies_view_is_unaffected(self):
        import dash_view
        v = dash_view.pod_view(self._two_room_batch(), "Techies")
        self.assertTrue(v["sessions"])
        self.assertTrue(all(s["pod"] == "Techies" for s in v["sessions"]))

    def test_a_batch_with_no_complement_offers_none(self):
        """The regression guard: B17-B34 have no PODs at all and must not grow
        a phantom Common view."""
        import dash_view
        import pods as _p
        pod_rows = [""] * 10
        rows = _pod_tab(pod_rows, [["Present"] * 6 + ["Absent"] * 4],
                        ["2026_08_15"])
        b = data.build_batch(rows, "B20", {})
        self.assertFalse([s for s in b["sessions"] if s.get("pod") == _p.COMMON])
        v = dash_view.pod_view(b, _p.COMMON)
        self.assertEqual(v["sessions"], [])


class TestDomainSplitInsideTheComplement(unittest.TestCase):
    """From B35 a batch's first two weekends run only Techies + everybody else,
    so Finance never gets a room of its own until week three and its view was
    simply empty for those dates. The roster still knows who is Finance, so the
    complement room's Finance share is recoverable."""

    def _batch(self):
        pod_rows = (["Techies"] * 4 + ["Finance"] * 3 + ["AI Generalist"] * 3)
        # Techies sat out the main room; of the others, 2 Finance + 1 Generalist came
        common = ["Absent"] * 4 + ["Present", "Present", "Absent"] + ["Present", "Absent", "Absent"]
        techies = ["Present"] * 3 + ["Absent"] + [""] * 6
        rows = _pod_tab(pod_rows, [common, techies],
                        ["2026_08_15", "2026_08_15 | Techies"])
        return data.build_batch(rows, "B35", {})

    def test_the_complement_session_carries_a_per_domain_split(self):
        b = self._batch()
        comp = next(s for s in b["sessions"] if s.get("pod") == "Common")
        self.assertEqual(comp["pod_split"]["Finance"], {"present": 2, "total": 3, "pct": 66.7})
        self.assertEqual(comp["pod_split"]["Generalist"]["total"], 3)
        self.assertNotIn("Techies", comp["pod_split"])   # they were not invited

    def test_the_domain_view_shows_its_share_of_that_session(self):
        import dash_view
        v = dash_view.pod_view(self._batch(), "Finance")
        self.assertEqual(len(v["sessions"]), 1)
        self.assertEqual(v["sessions"][0]["present"], 2)
        self.assertEqual(v["sessions"][0]["total"], 3)      # Finance, not the room's 6
        self.assertTrue(v["sessions"][0]["within_common"])

    def test_a_domain_with_its_own_room_that_day_is_not_counted_twice(self):
        """Techies met separately, so they are excluded from the complement and
        must appear exactly once."""
        import dash_view
        v = dash_view.pod_view(self._batch(), "Techies")
        self.assertEqual(len(v["sessions"]), 1)
        self.assertFalse(v["sessions"][0].get("within_common"))

    def test_an_all_domains_session_is_split_too(self):
        """A batch's first two or three sessions are All Domains - one room,
        everyone invited. They are exactly where a domain's attendance used to
        be invisible, so they get the same breakdown."""
        pod_rows = ["Finance"] * 4 + ["Data"] * 6
        whole = ["Present"] * 3 + ["Absent"] + ["Present"] * 2 + ["Absent"] * 4
        rows = _pod_tab(pod_rows, [whole], ["2026_08_15"])
        b = data.build_batch(rows, "B35", {})
        sx = next(s for s in b["sessions"] if s.get("mm"))
        self.assertFalse(sx.get("pod"))                       # still All Domains
        self.assertEqual(sx["pod_split"]["Finance"], {"present": 3, "total": 4, "pct": 75.0})
        self.assertEqual(sx["pod_split"]["Data"], {"present": 2, "total": 6, "pct": 33.3})

    def test_a_pod_with_its_own_room_is_left_out_of_the_all_domains_split(self):
        """Otherwise Techies show up twice for one date - once in their own
        room, once inside the All Domains breakdown."""
        pod_rows = ["Techies"] * 4 + ["Finance"] * 6
        whole = ["Present"] * 4 + ["Present"] * 3 + ["Absent"] * 3
        techies = ["Present"] * 4 + [""] * 6
        rows = _pod_tab(pod_rows, [whole, techies],
                        ["2026_08_15", "2026_08_15 | Techies"])
        b = data.build_batch(rows, "B35", {})
        whole_sx = next(s for s in b["sessions"] if s.get("mm") and not s.get("pod"))
        self.assertIn("Finance", whole_sx["pod_split"])
        self.assertNotIn("Techies", whole_sx["pod_split"])

    def test_a_single_pods_own_room_needs_no_split(self):
        pod_rows = ["Techies"] * 4 + ["Finance"] * 6
        techies = ["Present"] * 3 + ["Absent"] + [""] * 6
        rows = _pod_tab(pod_rows, [techies], ["2026_08_15 | Techies"])
        b = data.build_batch(rows, "B35", {})
        sx = next(s for s in b["sessions"] if s.get("mm"))
        self.assertEqual(sx["pod_split"], {})


class TestDomainMatrixHtml(unittest.TestCase):
    """The split is only useful if it is on screen. `domain_matrix_html` is pure
    so the static site can reuse it, and it must stay silent for the batches
    that have no domains at all rather than rendering an empty table."""

    def test_it_lays_out_domain_by_date(self):
        import dash_view
        pod_rows = ["Finance"] * 4 + ["Data"] * 6
        whole = ["Present"] * 3 + ["Absent"] + ["Present"] * 2 + ["Absent"] * 4
        rows = _pod_tab(pod_rows, [whole], ["2026_08_15"])
        b = data.build_batch(rows, "B35", {})
        html = dash_view.domain_matrix_html(b)
        self.assertIn("Finance", html)
        self.assertIn("Data", html)
        self.assertIn("3/4", html)        # Finance present/total
        self.assertIn("2/6", html)        # Data present/total

    def test_a_batch_with_no_domains_renders_nothing(self):
        import dash_view
        rows = _pod_tab([""] * 10, [["Present"] * 6 + ["Absent"] * 4],
                        ["2026_08_15"])
        b = data.build_batch(rows, "B20", {})
        self.assertEqual(dash_view.domain_matrix_html(b), "")

    def test_a_single_pods_own_room_renders_nothing(self):
        import dash_view
        pod_rows = ["Techies"] * 4 + ["Finance"] * 6
        rows = _pod_tab(pod_rows, [["Present"] * 3 + ["Absent"] + [""] * 6],
                        ["2026_08_15 | Techies"])
        b = data.build_batch(rows, "B35", {})
        self.assertEqual(dash_view.domain_matrix_html(b), "")


class TestSplitDropsNonBreakdowns(unittest.TestCase):
    """A single bucket covering everyone is not a breakdown - it restates the
    session's own total. B17-B34 hit this two ways: with no POD column every
    student keys on "", and with a column full of blanks they all key on
    "Unassigned". 464 of 491 stored splits were one such bucket."""

    def test_a_batch_with_no_real_pods_stores_no_split(self):
        rows = _pod_tab([""] * 10, [["Present"] * 6 + ["Absent"] * 4],
                        ["2026_08_15"])
        b = data.build_batch(rows, "B20", {})
        for s in b["sessions"]:
            if s.get("mm"):
                self.assertEqual(s["pod_split"], {}, s)

    def test_two_real_domains_are_kept(self):
        rows = _pod_tab(["Finance"] * 4 + ["Data"] * 6,
                        [["Present"] * 3 + ["Absent"] + ["Present"] * 2 + ["Absent"] * 4],
                        ["2026_08_15"])
        b = data.build_batch(rows, "B35", {})
        sx = next(s for s in b["sessions"] if s.get("mm"))
        self.assertEqual(sorted(sx["pod_split"]), ["Data", "Finance"])


class TestTwoSessionsInOneWebinar(unittest.TestCase):
    """One Zoom room, two sessions from different programmes.

    Owner's explanation, 2026-09-22: `AI CAP B40 - Common , BSIAI Accelerator B1`
    means that room hosted AI CAP B40's Common session AND BSIAI Accelerator B1's
    (B1 and B2 being different BSIAI batches). The label is deliberate.

    Reading only the tail after the last hyphen saw `Common , BSIAI Accelerator
    B1` as one domain name, found nothing, and warned "POD not recognised -
    counted against the whole batch. Add it to pods._ALIASES." on three real,
    correctly-labelled sessions every run - and every clause of that was wrong.
    """

    REAL = ["AI CAP B40 - Common , BSIAI Accelerator B1",
            "AI CAP B40 - Common , BSIAI Accelerator  B1",
            "AI CAP B41 - Common , BSIAI Accelerator B2"]

    def test_the_ai_cap_side_is_read_as_the_whole_batch_marker(self):
        for label in self.REAL:
            self.assertEqual(pods.from_l2_label(label), pods.WHOLE_BATCH, label)

    def test_it_no_longer_looks_unrecognised(self):
        # None is what made the pipeline warn.
        for label in self.REAL:
            self.assertIsNotNone(pods.from_l2_label(label), label)

    def test_both_programmes_still_come_out_of_the_same_cell(self):
        # The POD side answering for AI CAP must not change the batch side.
        import attendance_core as ac
        self.assertEqual(ac.extract_batches(self.REAL[0]),
                         {("CAP", 40), ("BSIAI", 1)})
        self.assertEqual(ac.extract_batches(self.REAL[2]),
                         {("CAP", 41), ("BSIAI", 2)})

    def test_a_named_domain_sharing_a_room_still_wins(self):
        self.assertEqual(
            pods.from_l2_label("AI CAP B40 - Techies , BSIAI Accelerator B1"),
            "Techies")

    def test_a_genuinely_unknown_tail_still_returns_None(self):
        # The warning must keep firing for something nobody has taught it.
        self.assertIsNone(pods.from_l2_label("AI CAP B35 - Nonsense"))
        self.assertIsNone(pods.from_l2_label("AI CAP B35 - Nonsense , Drivel"))


class TestCompoundAndDashlessLabels(unittest.TestCase):
    """Two L2 spellings from the 26-27 Sep 2026 weekend (B35-B38) that the
    label reader could not see, and that between them put four batches'
    weekend at 2-3%.

    26 Sep, L2 row 235:  'AI CAP B35, B36, B37, B38 Finance'  - no dash before
    the POD (every sibling row that weekend reads '..., B38 - Finance'). Read
    as a whole-batch session, a ~300-person Finance room was scored against
    3,260: B35 2.9%, B36 2.1%, B37 3.2%, B38 2.3%.

    27 Sep, L2 row 267:  'AI CAP B35 , B36 , B37 , B38 - S/M/HR + Content
    Creators'  - ONE room inviting TWO PODs. The tail resolved to nothing, the
    marker wrote an unlabelled column, and data.py re-read it as the day's
    Common complement (right number by luck, wrong name, and wrong the moment
    a third POD lacks a room).

    Measured over all 578 distinct Batch Name cells in L2's history, exactly
    these two labels resolve differently after the fix.
    """
    COMPOUND = "Sales/Marketing/HR + Content Creators"

    def test_a_dashless_trailing_pod_is_still_that_pod(self):
        self.assertEqual(pods.from_l2_label("AI CAP B35, B36, B37, B38 Finance"),
                         "Finance")

    def test_a_dashless_label_naming_no_pod_is_still_the_whole_batch(self):
        """Only a real POD name may change the answer: the batch token is
        stripped and what is left has to canon()."""
        for raw in ("AI CAP B35", "AI CAP B17 11AM", "AI CAP B17 + B21 11AM",
                    "AI CAP B39, B40, B41 All Domains ", "B37-42",
                    "ECAP B1 + B2 + B3 7:30 PM", "AI CAP 15+B17",
                    "B 22 IC + B 23,24 IC"):
            self.assertEqual(pods.from_l2_label(raw), pods.WHOLE_BATCH, raw)

    def test_a_plus_tail_is_one_canonical_compound_in_a_fixed_order(self):
        self.assertEqual(pods.from_l2_label(
            "AI CAP B35 , B36 , B37 , B38 - S/M/HR + Content Creators"),
            self.COMPOUND)
        # spelling AND order are normalised: the join key must be ONE string
        self.assertEqual(pods.from_l2_label(
            "AI CAP B35 - Content Creators + Sales/Marketing/HR"), self.COMPOUND)
        self.assertEqual(pods.from_l2_label("AICAPB35, B36-Finance+Data"),
                         "Finance + Data")

    def test_a_whole_batch_alias_inside_a_compound_is_the_whole_batch(self):
        self.assertEqual(pods.from_l2_label("AI CAP B35 - All Domains + Techies"),
                         pods.WHOLE_BATCH)

    def test_one_unknown_item_still_returns_none_so_the_caller_warns(self):
        self.assertIsNone(pods.from_l2_label("AI CAP B35 - Finance + Robotics"))

    def test_members_gives_the_pods_back_without_reparsing(self):
        self.assertEqual(pods.members(self.COMPOUND),
                         ("Sales/Marketing/HR", "Content Creators"))
        self.assertEqual(pods.members("Techies"), ("Techies",))
        self.assertEqual(pods.members(pods.COMMON), (pods.COMMON,))
        self.assertEqual(pods.members(pods.UNKNOWN), (pods.UNKNOWN,))
        self.assertEqual(pods.members(""), ())
        self.assertEqual(pods.members(None), ())

    def test_the_join_key_round_trips_through_the_column_header(self):
        """Whatever string the room carries must be the SAME string at the
        marker's header, `_col_pod` reading it back, and `session_key`."""
        import attendance_core as ac
        header = ac._col_header("2026_09_27", self.COMPOUND)
        self.assertEqual(header, "2026_09_27 | " + self.COMPOUND)
        self.assertEqual(ac._col_pod(header), self.COMPOUND)
        self.assertEqual(ac.session_key(header), ("2026_09_27", self.COMPOUND))

    def test_the_poll_lookup_files_the_room_under_the_same_key(self):
        """polls.lookup_by_session_rows keys on from_l2_label, so the compound
        string data.py reads off the column finds the rating."""
        import io
        import polls
        from openpyxl import Workbook
        wb = Workbook(); wb.remove(wb.active)
        ws = wb.create_sheet("Sep 2026")
        ws.append(["Date", "Webinar ID", "Batch Name", "Topic"])
        ws.append(["09/27/2026", "555", "AI CAP B35 , B36 - S/M/HR + Content Creators",
                   "Spy, Swipe & Ship"])
        buf = io.BytesIO(); wb.save(buf)
        rows = [("poll_555_2026_09_27.csv",
                 {"session": 4.5, "trainer": 4.6, "recommend": 4.4, "responses": 7})]
        got, _ = polls.lookup_by_session_rows(rows, buf.getvalue())
        self.assertIn(("B35", "09_27", self.COMPOUND), got)
        self.assertIn(("B36", "09_27", self.COMPOUND), got)


class TestCompoundRoomScoring(unittest.TestCase):
    """27 Sep in miniature: 3 S/M/HR, 2 Content Creators, 4 Techies, 3 Finance.
    The compound room marks its two PODs; Techies ran their own room."""
    COMPOUND = "Sales/Marketing/HR + Content Creators"

    def _rows(self, with_plain=False):
        pod_rows = (["Sales/Marketing/HR"] * 3 + ["Content Creators"] * 2
                    + ["Techies"] * 4 + ["Finance"] * 3)
        # 2 of 3 S/M/HR and 1 of 2 Content Creators came. A Techie (row 5) is
        # ALSO marked Present on this column - never invited, must not count.
        compound = (["Present", "Present", "Absent"] + ["Present", "Absent"]
                    + ["Present", "", "", ""] + [""] * 3)
        techies = [""] * 5 + ["Present", "Present", "Present", "Absent"] + [""] * 3
        cols = [compound, techies]
        headers = [f"2026_09_27 | {self.COMPOUND}", "2026_09_27 | Techies"]
        if with_plain:
            # an unlabelled room the same day: everyone marked, only Finance came
            cols.append(["Absent"] * 9 + ["Present"] * 3)
            headers.append("2026_09_27")
        return _pod_tab(pod_rows, cols, headers)

    def _build(self, **kw):
        return data.build_batch(self._rows(**kw), "B35", {("B35", "09_27"): "x"})

    def _room(self, b):
        return next(s for s in b["sessions"] if s["pod"] == self.COMPOUND)

    def test_the_denominator_is_the_sum_of_both_pods(self):
        s = self._room(self._build())
        self.assertEqual(s["total"], 5)          # 3 S/M/HR + 2 CC: not 12, not 3
        self.assertEqual(s["present"], 3)
        self.assertEqual(s["pct"], 60.0)
        self.assertEqual(s["absent"], 2)

    def test_a_student_in_either_pod_is_invited_and_a_third_pod_is_not(self):
        s = self._room(self._build())
        self.assertEqual(s["present"], 3)        # the Present Techie is not counted
        self.assertEqual(s["pod_split"]["Sales/Marketing/HR"],
                         {"present": 2, "total": 3, "pct": 66.7})
        self.assertEqual(s["pod_split"]["Content Creators"],
                         {"present": 1, "total": 2, "pct": 50.0})
        self.assertNotIn("Techies", s["pod_split"])

    def test_it_displays_as_the_compound_not_as_common(self):
        b = self._build()
        s = self._room(b)
        self.assertEqual(s["excl"], [])          # a named room, not a complement
        self.assertFalse([x for x in b["sessions"] if x["pod"] == pods.COMMON])

    def test_the_same_days_common_complement_excludes_both_pods(self):
        b = self._build(with_plain=True)
        comp = next(s for s in b["sessions"] if s["pod"] == pods.COMMON)
        self.assertEqual(comp["excl"],
                         ["Content Creators", "Sales/Marketing/HR", "Techies"])
        self.assertEqual(comp["total"], 3)       # Finance only
        self.assertEqual(comp["present"], 3)
        self.assertEqual(comp["pct"], 100.0)
        # and the compound room itself is untouched by the extra column
        self.assertEqual(self._room(b)["total"], 5)

    def test_each_member_pods_view_shows_its_share_of_the_room(self):
        import dash_view
        v = dash_view.pod_view(self._build(), "Content Creators")
        self.assertEqual(len(v["sessions"]), 1)
        self.assertEqual(v["sessions"][0]["present"], 1)
        self.assertEqual(v["sessions"][0]["total"], 2)   # CC, not the room's 5
        v = dash_view.pod_view(self._build(), "Techies")
        self.assertEqual(len(v["sessions"]), 1)          # their own room only
        self.assertEqual(v["sessions"][0]["total"], 4)

    def test_the_date_rollup_invites_both_pods_once_and_finance_not_at_all(self):
        d = self._build()["by_date"][0]
        self.assertEqual(d["total"], 9)          # 5 compound + 4 Techies
        self.assertEqual(d["present"], 6)        # 2 + 1 + 3
        self.assertEqual(d["n_pods"], 2)

    def test_closing_type_attendance_counts_the_room_for_both_pods(self):
        """Per-person denominators: S/M/HR and CC each have one session (the
        compound room), Techies one, Finance none -> 6 present of 9 slots."""
        b = self._build()
        self.assertEqual(len(b["closing"]), 1)
        self.assertEqual(b["closing"][0]["att"], round(6 / 9 * 100, 1))

    def test_the_payload_still_json_encodes(self):
        import json
        json.dumps(self._build(with_plain=True))


class TestMarkerWritesTheCompoundColumn(unittest.TestCase):
    """End to end through the marker: the L2 label, the header it writes, and
    who it marks. Neither the roster nor the folder spells the compound the
    way L2 does, so the column header is where the key is fixed."""
    COMPOUND = "Sales/Marketing/HR + Content Creators"
    SMHR = ("smhr@x.com", "919000000001")
    CC = ("cc@x.com", "919000000002")
    TECH = ("tech@x.com", "919000000003")
    FIN = ("fin@x.com", "919000000004")

    def _roster(self):
        import io
        from openpyxl import Workbook
        wb = Workbook(); wb.remove(wb.active)
        ws = wb.create_sheet("AI CAP B35")
        ws.append(["Country", "Registered Number", "Registered Mail", "WhatsApp",
                   "Broadcast", "Batch", "Amount", "Payment", "Close Type",
                   "POD Prefrence"])
        for (em, ph), pod in ((self.SMHR, "Sales/Marketing/HR"),
                              (self.CC, "Content Creators - AI Career Accelerator Program B35"),
                              (self.TECH, "Techies"), (self.FIN, "Finance")):
            ws.append([91, ph, em, "", "", "B35", 0, "Full Paid", "BDA Closing", pod])
        buf = io.BytesIO(); wb.save(buf); return buf.getvalue()

    def _l2(self, label, wid, date="09/27/2026"):
        import io
        from openpyxl import Workbook
        wb = Workbook(); wb.remove(wb.active)
        ws = wb.create_sheet("Sep 2026")
        ws.append(["Date", "Webinar ID", "Batch Name", "Topic"])
        ws.append([date, wid, label, "Session"])
        buf = io.BytesIO(); wb.save(buf); return buf.getvalue()

    def _run(self, label, folder, wid, ymd, attendees):
        import attendance_core as ac
        from tests.test_cross_room import report_csv
        files = [(f"{folder}/attendee_{wid}_{ymd}.csv", report_csv(attendees))]
        mdy = f"{ymd[5:7]}/{ymd[8:10]}/{ymd[:4]}"          # L2 writes 09/27/2026
        out, report, warns = ac.process_files(
            self._roster(), self._l2(label, wid, mdy), files, values_only=True)
        return out, report, warns

    def _headers(self, out):
        import io
        from openpyxl import load_workbook
        ws = load_workbook(io.BytesIO(out))["AI CAP B35"]
        return [str(ws.cell(1, c).value) for c in range(11, ws.max_column + 1)]

    def test_the_plus_label_marks_both_pods_under_one_compound_header(self):
        out, report, warns = self._run(
            "AI CAP B35 , B36 , B37 , B38 - S/M/HR + Content Creators",
            "2026-09-27 - AI CAP B35 , B36 , B37 , B38 - S M HR + Content Creators - Spy Swipe Ship",
            "99900000027", "2026_09_27", [self.SMHR, self.CC, self.TECH])
        self.assertEqual(self._headers(out), ["2026_09_27 | " + self.COMPOUND])
        r = report[0]
        self.assertEqual(r["pod"], self.COMPOUND)
        self.assertEqual(r["present"], 2, "S/M/HR + Content Creators")
        self.assertEqual(r["total"], 2, "denominator is the two PODs, not the batch")
        self.assertEqual(r["outside"], 1, "the Techie is reported, not counted")
        self.assertFalse([w for w in warns if "not recognised" in w], warns)

    def test_the_dashless_label_marks_a_finance_column(self):
        out, report, warns = self._run(
            "AI CAP B35, B36, B37, B38 Finance",
            "2026-09-26 - AI CAP B35, B36, B37, B38 Finance - Forecasting",
            "99900000026", "2026_09_26", [self.FIN, self.TECH])
        self.assertEqual(self._headers(out), ["2026_09_26 | Finance"])
        r = report[0]
        self.assertEqual(r["pod"], "Finance")
        self.assertEqual((r["present"], r["total"]), (1, 1))   # NOT 2 of 4
        self.assertEqual(r["outside"], 1)


class TestCompoundRoomWithOneMemberAbsentFromTheBatch(unittest.TestCase):
    """A batch with 3 S/M/HR, 4 Techies and NO Content Creators sits in the
    compound room. Its breakdown collapses - `data._real_split` drops a lone
    bucket covering everyone invited - so the S/M/HR filter has no share to
    take from it. The room IS that POD's session (its denominator is already
    S/M/HR's strength), and the filter must show it whole rather than show
    nothing for the date. Not live for 26-27 Sep 2026 (every batch has both
    PODs) but it is the mechanism the compound room introduced."""
    COMPOUND = "Sales/Marketing/HR + Content Creators"

    def _build(self):
        pod_rows = ["Sales/Marketing/HR"] * 3 + ["Techies"] * 4
        compound = ["Present", "Present", "Absent"] + [""] * 4
        techies = [""] * 3 + ["Present", "Present", "Present", "Absent"]
        rows = _pod_tab(pod_rows, [compound, techies],
                        [f"2026_09_27 | {self.COMPOUND}", "2026_09_27 | Techies"])
        return data.build_batch(rows, "B38", {("B38", "09_27"): "x"})

    def test_the_room_scores_against_the_one_pod_that_is_present(self):
        s = next(x for x in self._build()["sessions"] if x["pod"] == self.COMPOUND)
        self.assertEqual((s["present"], s["total"]), (2, 3))
        self.assertEqual(s["pod_split"], {})     # one bucket is not a breakdown

    def test_the_present_members_filter_shows_the_room_whole(self):
        import dash_view
        v = dash_view.pod_view(self._build(), "Sales/Marketing/HR")
        self.assertEqual([(x["present"], x["total"], x["pod"]) for x in v["sessions"]],
                         [(2, 3, "Sales/Marketing/HR")])
        self.assertEqual(v["strength"], 3)
        self.assertEqual(v["avg_pct"], 66.7)

    def test_the_absent_members_filter_shows_nothing_and_is_not_offered(self):
        import dash_view
        b = self._build()
        self.assertNotIn("Content Creators", dash_view.pod_names(b))
        self.assertEqual(dash_view.pod_view(b, "Content Creators")["sessions"], [])

    def test_a_third_pod_does_not_inherit_the_room(self):
        import dash_view
        v = dash_view.pod_view(self._build(), "Techies")
        self.assertEqual([(x["present"], x["total"]) for x in v["sessions"]], [(3, 4)])

    def test_a_room_with_a_real_breakdown_still_hands_out_shares_not_the_whole(self):
        """With both PODs present the existing path - a share per member - is
        the one taken, never the whole-row fallback."""
        import dash_view
        pod_rows = ["Sales/Marketing/HR"] * 3 + ["Content Creators"] * 2
        compound = ["Present", "Present", "Absent", "Present", "Absent"]
        rows = _pod_tab(pod_rows, [compound], [f"2026_09_27 | {self.COMPOUND}"])
        b = data.build_batch(rows, "B38", {("B38", "09_27"): "x"})
        v = dash_view.pod_view(b, "Content Creators")
        self.assertEqual([(x["present"], x["total"]) for x in v["sessions"]], [(1, 2)])
        self.assertTrue(v["sessions"][0].get("within_common"))

    def test_every_row_still_renders(self):
        import dash_view
        b = self._build()
        for sel in dash_view.pod_names(b):
            dash_view.sessions_table_html(dash_view.pod_view(b, sel))


class TestFolderKeyMatchesTheMarkersHeader(unittest.TestCase):
    """The freeze: `attendance_core.folder_keys` (what the incremental skip
    compares) must contain the `session_key` of the header the marker wrote
    from L2, or the folder is downloaded and re-marked on every run. Both
    26-27 Sep 2026 folders failed this until `from_folder` learned the '+'
    and dashless spellings; the compound comes back as the SAME string L2's
    label resolves to, whether the folder spaces or slashes S/M/HR."""
    COMPOUND = "Sales/Marketing/HR + Content Creators"
    PLUS_L2 = "AI CAP B35 , B36 , B37 , B38 - S/M/HR + Content Creators"
    PLUS_FOLDER = ("2026-09-27 - AI CAP B35 , B36 , B37 , B38 - S M HR + "
                   "Content Creators - Spy Swipe Ship")
    DASHLESS_L2 = "AI CAP B35, B36, B37, B38 Finance"
    DASHLESS_FOLDER = "2026-09-26 - AI CAP B35, B36, B37, B38 Finance - Forecasting"

    def test_the_two_real_folders_resolve_to_the_columns_l2_names(self):
        self.assertEqual(pods.from_folder(self.PLUS_FOLDER), self.COMPOUND)
        self.assertEqual(pods.from_folder(
            self.PLUS_FOLDER.replace("S M HR", "S/M/HR")), self.COMPOUND)
        self.assertEqual(pods.from_folder(self.DASHLESS_FOLDER), "Finance")

    def test_folder_keys_contain_the_header_the_marker_writes(self):
        import attendance_core as ac
        for label, folder, ymd in ((self.PLUS_L2, self.PLUS_FOLDER, "2026_09_27"),
                                   (self.DASHLESS_L2, self.DASHLESS_FOLDER, "2026_09_26")):
            header = ac._col_header(ymd, pods.from_l2_label(label))
            self.assertIn(ac.session_key(header), ac.folder_keys(folder), folder)

    def test_the_older_spellings_still_read_the_same(self):
        for folder, want in (
            ("2026-08-23 - AI CAP B35 - Techies - Python with AI", "Techies"),
            ("2026-09-06 - AI CAP B37 8PM-Techis - Topic", "Techies"),
            ("2026-08-30 - AICAPB35, B36-Generalist - Topic", "Generalist"),
            ("2026-08-02 - AI CAP B33 - Office Productivity", None),
            ("2026-09-13 - AI CAP B39, B40, B41 All Domains - Topic", None),
            ("2026-09-27 - AI CAP B42 Day 2 - Topic", None),
            ("2026-09-27 - AI CAP B37 8PM - Topic", None),
            ("2026-09-27 - AI CAP B8 + B22 - Topic", None),
        ):
            self.assertEqual(pods.from_folder(folder), want, folder)

    def test_a_topic_ending_in_a_domain_word_is_not_a_domain(self):
        # the dashless rule runs only inside a segment that names a batch
        self.assertIsNone(pods.from_folder("2026-09-27 - AI CAP B35 - Session 2 Finance"))
        self.assertIsNone(pods.from_folder("2026-09-27 - AI CAP B35 - Excel + AI"))

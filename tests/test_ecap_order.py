"""ECAP must be grafted BEFORE the carry-forward: pipeline [1c] before [1d].

`carryforward.merge_marks` keeps a tab's history only when THIS WEEK'S roster
has that tab. Until 2026-09-30 the ECAP graft ran after the carry, so at merge
time the workbook held only CAP tabs, every ECAP column was dropped with the
warning "batch tab(s) in last week's workbook but NOT in this week's roster",
and the graft then re-added ECAP as a pristine roster - re-fetched and
re-marked from Drive on every run instead of frozen like every other batch,
and lost for good on any run that could not re-fetch it.

Both halves are pinned: the order that keeps ECAP's history, the order that
loses it (so nobody "simplifies" the graft back down), and the pipeline source
itself calling them in the right order.
"""
import io
import os
import sys
import unittest

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)

from openpyxl import load_workbook, Workbook            # noqa: E402

import carryforward                                     # noqa: E402
import ecap                                             # noqa: E402

FIXED = ["Contry code", "Registered Number", "Registered mail", "Contry code",
         "Whatsaap Number", "broadcast mail", "batch name", "amount",
         "Payment", "Closing Type", "POD pref"]


def tab(sheet, students, session_headers=()):
    """students: [(mail, phone, [marks per session header])]"""
    rows = [list(FIXED) + list(session_headers)]
    for mail, phone, marks in students:
        rows.append([91, phone, mail, 91, phone, mail, sheet, 1000,
                     "Full Paid", "BDA Closing", ""] + list(marks))
    return sheet, rows


def book(*tabs) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    for name, rows in tabs:
        ws = wb.create_sheet(name)
        for r in rows:
            ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def headers(b: bytes, sheet: str) -> list:
    wb = load_workbook(io.BytesIO(b), read_only=True, data_only=True)
    try:
        return [c for c in next(wb[sheet].iter_rows(min_row=1, max_row=1, values_only=True))]
    finally:
        wb.close()


class TestGraftBeforeCarry(unittest.TestCase):
    def setUp(self):
        # Last week's marked workbook: a CAP tab and an ECAP tab, each with one
        # marked session. This week's roster export: CAP only (the Sheet has no
        # ECAP tab; ECAP comes from the snapshot workbook, `extra`).
        self.prev = book(tab("AI CAP B29", [("a@x.com", "919000000001", ["Present"])], ["2026_05_16"]),
                         tab("AI ECAP B1", [("e@x.com", "919000000002", ["Present"])], ["2026_05_16"]))
        self.fresh = book(tab("AI CAP B29", [("a@x.com", "919000000001", [])]))
        self.extra = book(tab("AI ECAP B1", [("e@x.com", "919000000002", [])]))

    def test_graft_then_carry_keeps_ecap_history(self):
        grafted, rep = ecap.graft(self.fresh, self.extra)
        self.assertEqual(rep["added"], ["AI ECAP B1"])
        out, carry = carryforward.merge_marks(grafted, self.prev)
        self.assertIn("2026_05_16", headers(out, "AI ECAP B1"))
        self.assertIn("2026_05_16", headers(out, "AI CAP B29"))
        self.assertEqual(carry["carried"], 2)
        self.assertFalse([w for w in carry.get("warnings") or () if "NOT in this week" in w])

    def test_carry_then_graft_loses_it__the_bug_the_order_fixes(self):
        out, carry = carryforward.merge_marks(self.fresh, self.prev)
        self.assertTrue(any("ECAP" in w for w in carry.get("warnings") or ()),
                        "the carry must at least SAY it dropped ECAP")
        grafted, _ = ecap.graft(out, self.extra)
        self.assertNotIn("2026_05_16", headers(grafted, "AI ECAP B1"))

    def test_the_pipeline_grafts_before_it_carries(self):
        with open(os.path.join(ROOT, "pipeline.py"), encoding="utf-8") as fh:
            src = fh.read()
        graft_at = src.index('print("[1c] Grafting the ECAP roster tabs')
        carry_at = src.index('print("[1d] Carrying last week')
        self.assertLess(graft_at, carry_at)
        self.assertLess(src.index("ecap.graft(roster_bytes, _eb)"),
                        src.index("carryforward.merge_marks(roster_bytes, prev_marked)"))
        self.assertLess(carry_at, src.index("[2/8] Fetching attendee reports"))


if __name__ == "__main__":
    unittest.main()

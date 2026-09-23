"""ECAP: a second programme sharing L2 and the Drives with AI CAP.

Two things have to hold, and they pull in opposite directions.

**ECAP must reach the dashboard at all.** Every batch filter in the stack was
written as "AI CAP B<n>", so an `AI ECAP B1` tab was silently dropped - marked
into the workbook by a marker that understood it perfectly, then thrown away by
`data.batch_label` returning None.

**ECAP must never be mistaken for CAP.** The two programmes number their
cohorts independently: ECAP B1 is 206 people, CAP B1 is 3,985. A bare `B1` code
would collide, and `batch_key` sorting on the number alone would file ECAP B1
next to CAP B1 as though they were the same cohort a week apart.
"""
import io
import os
import sys
import unittest

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import attendance_core as ac  # noqa: E402
import dashboard_core as dc  # noqa: E402
import data as ddata  # noqa: E402
import ecap  # noqa: E402

HDR = ["Contry code", "Registered Number", "Registered mail", "Contry code",
       "Whatsaap Number", "broadcast mail", "batch name", "amount",
       "Payment", "Closing Type"]


def book(tabs):
    """{tab name: [rows]} -> xlsx bytes."""
    from openpyxl import Workbook
    wb = Workbook()
    wb.remove(wb.active)
    for name, rows in tabs.items():
        ws = wb.create_sheet(name)
        for r in rows:
            ws.append(list(r))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def roster_rows(batch, n=2, marks=None):
    rows = [list(HDR) + (["Topic"] if marks else [])]
    for i in range(n):
        rows.append([91, f"9190000000{i}", f"{batch.replace(' ', '')}_{i}@x.com",
                     "", "", "", batch, "", "Full Paid", "BDA Collection"]
                    + ([marks[i]] if marks else []))
    return rows


class TestTabDetection(unittest.TestCase):
    def test_the_real_tab_names_are_recognised(self):
        for n in ("AI ECAP B1", "AI ECAP B2", "AI ECAP B3", "ECAP B3",
                  "AI E-CAP B1", " AI ECAP B10 "):
            self.assertTrue(ecap.is_batch_tab(n), n)

    def test_cap_tabs_and_notes_are_not(self):
        # 'Read me' is a real tab in the supplied workbook, and grafting it
        # would put a page of prose in the dashboard as a batch.
        for n in ("Read me", "AI CAP B29", "AI CAP B3", "ECAP notes",
                  "Sheet1", "", None):
            self.assertFalse(ecap.is_batch_tab(n), repr(n))


class TestGraft(unittest.TestCase):
    def _base(self):
        return book({"AI CAP B29": roster_rows("AI CAP B29")})

    def _extra(self):
        return book({"Read me": [["ECAP B1-B3 in the roster format."]],
                     "AI ECAP B1": roster_rows("AI ECAP B1", 3),
                     "AI ECAP B2": roster_rows("AI ECAP B2", 2)})

    def _tabs(self, b):
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(b), read_only=True)
        try:
            return list(wb.sheetnames)
        finally:
            wb.close()

    def test_the_ecap_tabs_are_copied_and_the_prose_tab_is_not(self):
        out, rep = ecap.graft(self._base(), self._extra())
        self.assertEqual(self._tabs(out), ["AI CAP B29", "AI ECAP B1", "AI ECAP B2"])
        self.assertEqual(rep["added"], ["AI ECAP B1", "AI ECAP B2"])
        self.assertEqual(rep["rows"], 5, "3 + 2 students, headers not counted")

    def test_a_batch_ALREADY_in_the_base_is_skipped_not_overwritten(self):
        # The freeze depends on this. From the second run on, the base IS last
        # week's marked workbook: copying the pristine snapshot over it would
        # wipe every carried Present/Absent.
        base = book({"AI CAP B29": roster_rows("AI CAP B29"),
                     "AI ECAP B1": roster_rows("AI ECAP B1", 3, marks=["Present", "Absent", "Present"])})
        out, rep = ecap.graft(base, self._extra())
        self.assertEqual(rep["skipped"], ["AI ECAP B1"])
        self.assertEqual(rep["added"], ["AI ECAP B2"])
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(out))
        marks = [wb["AI ECAP B1"].cell(r, 11).value for r in (2, 3, 4)]
        wb.close()
        self.assertEqual(marks, ["Present", "Absent", "Present"],
                         "the carried marks survived the graft")

    def test_no_ecap_workbook_is_a_no_op_not_a_crash(self):
        base = self._base()
        out, rep = ecap.graft(base, b"")
        self.assertIs(out, base)
        self.assertEqual(rep["added"], [])

    def test_a_workbook_with_no_ecap_tab_says_so(self):
        _out, rep = ecap.graft(self._base(), book({"Read me": [["nothing here"]]}))
        self.assertEqual(rep["added"], [])
        self.assertTrue(rep["warnings"], "silence here looks identical to success")


class TestBatchCodes(unittest.TestCase):
    def test_cap_labels_are_unchanged(self):
        # Every batch published before ECAP existed must keep its name, or the
        # store's grid tables and GATE 6's comparison both break.
        self.assertEqual(ddata.batch_label("AI CAP B29"), "B29")
        self.assertEqual(ddata.batch_label("AI CAP B7"), "B7")
        self.assertEqual(ddata.batch_label("AICAP B17"), "B17")

    def test_ecap_labels_carry_the_programme(self):
        self.assertEqual(ddata.batch_label("AI ECAP B1"), "ECAP B1")
        self.assertEqual(ddata.batch_label("AI E-CAP B2"), "ECAP B2")

    def test_helper_tabs_are_still_ignored(self):
        self.assertIsNone(ddata.batch_label("B17 Att"))
        self.assertIsNone(ddata.batch_label("Auto pay"))

    def test_dashboard_core_agrees_with_data(self):
        # These two name the same thing in different modules. When they drift,
        # the store writes grid_<one> and the dashboard looks up <the other>.
        for tab in ("AI CAP B29", "AI CAP B7", "AI ECAP B1", "AI ECAP B3"):
            self.assertEqual(dc._clean_batch_name(tab), ddata.batch_label(tab), tab)

    def test_the_marker_keys_ecap_to_its_own_programme(self):
        # The guard that stops an ECAP tab being marked as a CAP batch.
        self.assertEqual(ac._sheet_key("AI ECAP B1"), ("ECAP", 1))
        self.assertEqual(ac._sheet_key("AI CAP B1"), ("CAP", 1))
        self.assertNotEqual(ac._sheet_key("AI ECAP B1"), ac._sheet_key("AI CAP B1"))


class TestOrdering(unittest.TestCase):
    def test_ecap_sorts_after_every_cap_batch(self):
        got = sorted(["ECAP B1", "B41", "B7", "ECAP B3", "B17"], key=dc.batch_key)
        self.assertEqual(got, ["B7", "B17", "B41", "ECAP B1", "ECAP B3"])

    def test_batch_key_still_returns_an_int(self):
        # pandas sort_values(key=...) maps this over a Series; a tuple would
        # change the dtype and sort differently.
        self.assertIsInstance(dc.batch_key("ECAP B1"), int)
        self.assertIsInstance(dc.batch_key("B29"), int)

    def test_data_build_puts_ecap_last(self):
        sheets = {"AI ECAP B1": roster_rows("AI ECAP B1", 2, marks=["Present", "Absent"]),
                  "AI CAP B29": roster_rows("AI CAP B29", 2, marks=["Present", "Present"])}
        DATA, _summary = ddata.build(sheets)
        self.assertEqual(list(DATA), ["B29", "ECAP B1"])


class TestTopicAndTrainerLookup(unittest.TestCase):
    """The programme has to be part of the topic/label/mentor key.

    It was not, and the failure was almost invisible: the date-only fallback
    still found A topic, so `REQUIRE_L2` kept every ECAP session and the page
    looked full - of whichever CAP session shared that date. 62 of 72 ECAP
    sessions carried another batch's title on the first build. Only the Trainer
    column told the truth, because the mentor has no date-only fallback.
    """

    def _l2(self):
        from openpyxl import Workbook
        wb = Workbook()
        wb.remove(wb.active)
        ws = wb.create_sheet("Sep 2026")
        ws.append(["Date", "Webinar ID", "Batch Name", "Topic Name", "Mentor"])
        ws.append(["09/12/2026", "81489430961", "ECAP B1", "Cursor AI", "Rajat"])
        ws.append(["09/12/2026", "91695866411", "AI CAP B29", "RAG beginner", "Amit"])
        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()

    def _files(self):
        return [("2026-09-12 - ECAP B1/attendee_81489430961_2026_09_12.csv", b""),
                ("2026-09-12 - AI CAP B29/attendee_91695866411_2026_09_12.csv", b"")]

    def test_ecap_is_keyed_under_its_own_programme(self):
        import sheets as dsheets
        lookup, _labels, mentors = dsheets.webinar_topic_lookup(
            self._files(), self._l2(), with_labels=True, with_mentors=True)
        self.assertEqual(lookup.get(("ECAP B1", "09_12", "")), "Cursor AI")
        self.assertEqual(mentors.get(("ECAP B1", "09_12", "")), "Rajat")

    def test_it_does_NOT_leak_onto_cap_B1(self):
        # The original bug: ECAP B1's row filed itself under "B1", which is AI
        # CAP B1 - a different programme's 3,985-person cohort.
        import sheets as dsheets
        lookup, _labels, mentors = dsheets.webinar_topic_lookup(
            self._files(), self._l2(), with_labels=True, with_mentors=True)
        self.assertNotIn(("B1", "09_12", ""), lookup)
        self.assertNotIn(("B1", "09_12", ""), mentors)

    def test_cap_keys_are_unchanged(self):
        import sheets as dsheets
        lookup, _labels, mentors = dsheets.webinar_topic_lookup(
            self._files(), self._l2(), with_labels=True, with_mentors=True)
        self.assertEqual(lookup.get(("B29", "09_12", "")), "RAG beginner")
        self.assertEqual(mentors.get(("B29", "09_12", "")), "Amit")

    def test_the_key_matches_what_data_looks_up_with(self):
        # The two sides of the same join, compared directly - this is the
        # assertion that would have failed on the first ECAP build.
        import sheets as dsheets
        lookup = dsheets.webinar_topic_lookup(self._files(), self._l2())
        for tab in ("AI ECAP B1", "AI CAP B29"):
            code = ddata.batch_label(tab)
            self.assertIn((code, "09_12", ""), lookup, tab)


if __name__ == "__main__":
    unittest.main()

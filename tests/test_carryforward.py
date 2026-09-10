"""Carrying last week's marks onto this week's roster.

The whole freeze design rests on this module: with no weekly rebuild, a column
it drops is history gone, and a mark it invents is history that never happened.
So these tests pin both directions, and they pin the two rules that make the
merge faithful to what a full re-mark would have produced - identity by
email-or-phone, and 'Absent' only for students who were in scope.
"""
import io
import os
import sys
import unittest

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from openpyxl import load_workbook, Workbook            # noqa: E402

import attendance_core as ac                            # noqa: E402
import carryforward as cf                               # noqa: E402

# Ten fixed columns, exactly like the real roster, then POD pref, then sessions.
FIXED = ["Contry code", "Registered Number", "Registered mail", "Contry code",
         "Whatsaap Number", "broadcast mail", "batch name", "amount",
         "Payment", "Closing Type", "POD pref"]


def tab(students, session_headers=(), sheet="AI CAP B35"):
    """students: [(mail, phone, pod, [marks per session_header])]"""
    rows = [list(FIXED) + list(session_headers)]
    for mail, phone, pod, marks in students:
        rows.append([91, phone, mail, 91, phone, mail, sheet, 1000,
                     "Full Paid", "BDA Closing", pod] + list(marks))
    return sheet, rows


def book(*tabs) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    for name, rows in tabs:
        ws = wb.create_sheet(name)
        for r in rows:
            ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def read(b: bytes, sheet="AI CAP B35"):
    """-> (header row 1 as list, {registered mail: {session header: value}})"""
    wb = load_workbook(io.BytesIO(b), data_only=True)
    ws = wb[sheet]
    hdr = [c.value for c in ws[1]]
    out = {}
    for r in range(2, (ws.max_row or 1) + 1):
        mail = ac._cell_email(ws.cell(r, 3).value)
        if not mail:
            continue
        out[mail] = {hdr[c - 1]: ws.cell(r, c).value
                     for c in range(1, len(hdr) + 1)
                     if hdr[c - 1] and ac._mmdd(hdr[c - 1])}
    wb.close()
    return hdr, out


class TestPassthrough(unittest.TestCase):
    def test_no_previous_workbook_is_a_full_run(self):
        fresh = book(tab([("a@x.com", "919000000001", "", [])]))
        out, rep = cf.merge_marks(fresh, None)
        self.assertIs(out, fresh)                 # untouched, not even re-saved
        self.assertEqual(rep["carried"], 0)
        self.assertTrue(rep["warnings"])

    def test_a_previous_workbook_with_no_session_columns_is_a_full_run(self):
        fresh = book(tab([("a@x.com", "919000000001", "", [])]))
        prev = book(tab([("a@x.com", "919000000001", "", [])]))
        out, rep = cf.merge_marks(fresh, prev)
        self.assertIs(out, fresh)
        self.assertEqual(rep["carried"], 0)

    def test_a_brand_new_batch_tab_passes_through_untouched(self):
        prev = book(tab([("a@x.com", "919000000001", "", ["Present"])],
                        ["2026_08_23"]))
        fresh = book(tab([("a@x.com", "919000000001", "", [])], []),
                     tab([("n@x.com", "919000000009", "", [])], [],
                         sheet="AI CAP B40"))
        out, rep = cf.merge_marks(fresh, prev)
        hdr40, rows40 = read(out, "AI CAP B40")
        self.assertEqual([h for h in hdr40 if h and ac._mmdd(h)], [])
        self.assertEqual(rows40["n@x.com"], {})
        # ...while B35 did get its column
        hdr35, _ = read(out)
        self.assertIn("2026_08_23", hdr35)
        self.assertEqual(rep["tabs"], 1)


class TestCarrying(unittest.TestCase):
    def test_columns_are_carried_with_their_header_verbatim(self):
        prev = book(tab([("a@x.com", "919000000001", "Techies", ["Present", "Absent"]),
                         ("b@x.com", "919000000002", "Finance", ["", "Present"])],
                        ["2026_08_23 | Techies", "9th May"]))
        fresh = book(tab([("a@x.com", "919000000001", "Techies", []),
                          ("b@x.com", "919000000002", "Finance", [])]))
        out, rep = cf.merge_marks(fresh, prev)
        hdr, rows = read(out)
        # the year AND the pod survive, and a legacy header survives verbatim
        self.assertIn("2026_08_23 | Techies", hdr)
        self.assertIn("9th May", hdr)
        self.assertEqual(rows["a@x.com"]["2026_08_23 | Techies"], "Present")
        self.assertEqual(rows["a@x.com"]["9th May"], "Absent")
        self.assertIsNone(rows["b@x.com"]["2026_08_23 | Techies"])   # blank stays blank
        self.assertEqual(rows["b@x.com"]["9th May"], "Present")
        self.assertEqual(rep["carried"], 2)
        self.assertEqual(rep["matched"], 2)

    def test_marks_follow_the_student_when_rows_move(self):
        prev = book(tab([("a@x.com", "919000000001", "", ["Present"]),
                         ("b@x.com", "919000000002", "", ["Absent"])],
                        ["2026_08_23"]))
        # reversed, with a newcomer wedged in the middle
        fresh = book(tab([("b@x.com", "919000000002", "", []),
                          ("new@x.com", "919000000003", "", []),
                          ("a@x.com", "919000000001", "", [])]))
        _hdr, rows = read(cf.merge_marks(fresh, prev)[0])
        self.assertEqual(rows["a@x.com"]["2026_08_23"], "Present")
        self.assertEqual(rows["b@x.com"]["2026_08_23"], "Absent")

    def test_a_changed_email_still_matches_on_the_last_ten_phone_digits(self):
        # The roster is 12-digit (91 prefix) and gets re-typed; a re-mark would
        # still have found this student, so the carry must too.
        prev = book(tab([("old@x.com", "919886440098", "", ["Present"])],
                        ["2026_08_23"]))
        fresh = book(tab([("new@x.com", "9886440098", "", [])]))
        out, rep = cf.merge_marks(fresh, prev)
        _hdr, rows = read(out)
        self.assertEqual(rows["new@x.com"]["2026_08_23"], "Present")
        self.assertEqual(rep["matched"], 1)
        self.assertEqual(rep["new_students"], 0)

    def test_a_departed_student_simply_disappears(self):
        prev = book(tab([("a@x.com", "919000000001", "", ["Present"]),
                         ("gone@x.com", "919000000002", "", ["Present"])],
                        ["2026_08_23"]))
        fresh = book(tab([("a@x.com", "919000000001", "", [])]))
        _hdr, rows = read(cf.merge_marks(fresh, prev)[0])
        self.assertEqual(set(rows), {"a@x.com"})

    def test_a_column_the_export_already_has_is_written_onto_not_duplicated(self):
        """B17-B28 keep their Present/Absent columns in the roster Sheet itself,
        and LAST WEEK'S VALUE WINS there.

        Measured on the real corpus: a full run re-marks two of those legacy
        columns as a side effect of the shared 'AI CAP B17 + B21' folders, so
        last week's workbook holds what the dashboard has been showing. Leaving
        the Sheet's own value in place instead moved 82 cells Present -> Absent
        on two of B17's published sessions - the first thing GATE 6 would have
        refused.
        """
        prev = book(tab([("a@x.com", "919000000001", "", ["Present"])],
                        ["2026_08_23"]))
        fresh = book(tab([("a@x.com", "919000000001", "", ["Absent"])],
                         ["2026_08_23"]))
        out, rep = cf.merge_marks(fresh, prev)
        hdr, rows = read(out)
        self.assertEqual([h for h in hdr if h and ac.session_key(h)],
                         ["2026_08_23"])          # one column, not two
        self.assertEqual(rows["a@x.com"]["2026_08_23"], "Present")
        self.assertEqual(rep["carried"], 0)
        self.assertEqual(rep["already_present"], 1)
        self.assertEqual(rep["overridden"], 1)    # the disagreement is counted

    def test_an_agreeing_cell_is_not_counted_as_an_override(self):
        prev = book(tab([("a@x.com", "919000000001", "", ["Present"])],
                        ["2026_08_23"]))
        fresh = book(tab([("a@x.com", "919000000001", "", ["Present"])],
                         ["2026_08_23"]))
        _out, rep = cf.merge_marks(fresh, prev)
        self.assertEqual(rep["overridden"], 0)


class TestNewStudents(unittest.TestCase):
    def test_a_new_student_is_absent_for_a_whole_batch_session(self):
        # What a full re-mark writes for them, so the numbers do not move when
        # the mode changes.
        prev = book(tab([("a@x.com", "919000000001", "", ["Present"])],
                        ["2026_08_23"]))
        fresh = book(tab([("a@x.com", "919000000001", "", []),
                          ("new@x.com", "919000000009", "", [])]))
        out, rep = cf.merge_marks(fresh, prev)
        _hdr, rows = read(out)
        self.assertEqual(rows["new@x.com"]["2026_08_23"], "Absent")
        self.assertEqual(rep["new_students"], 1)

    def test_a_new_student_is_blank_for_a_pod_they_are_not_in(self):
        """The bug this rule exists to prevent.

        A POD session was only offered to that POD - the marker leaves everyone
        else empty. Filling 'Absent' regardless would put ten false absences
        against every student on a day that ran eleven PODs.
        """
        prev = book(tab([("a@x.com", "919000000001", "Techies",
                          ["Present", "Present"])],
                        ["2026_08_23 | Techies", "2026_08_23 | Finance"]))
        fresh = book(tab([("a@x.com", "919000000001", "Techies", []),
                          ("new@x.com", "919000000009", "Finance", [])]))
        _hdr, rows = read(cf.merge_marks(fresh, prev)[0])
        self.assertEqual(rows["new@x.com"]["2026_08_23 | Finance"], "Absent")
        self.assertIsNone(rows["new@x.com"]["2026_08_23 | Techies"])

    def test_a_pod_cell_carrying_the_programme_suffix_still_resolves(self):
        # The header says 'Techies'; the cell says
        # 'Techies - AI Career Accelerator Program B35'. Comparing raw matched
        # ZERO of 3,236 students the last time somebody tried it.
        prev = book(tab([("a@x.com", "919000000001", "Techies", ["Present"])],
                        ["2026_08_23 | Techies"]))
        fresh = book(tab([
            ("a@x.com", "919000000001", "Techies", []),
            ("new@x.com", "919000000009",
             "Techies - AI Career Accelerator Program B35", [])]))
        _hdr, rows = read(cf.merge_marks(fresh, prev)[0])
        self.assertEqual(rows["new@x.com"]["2026_08_23 | Techies"], "Absent")


class TestExistingSessions(unittest.TestCase):
    """What counts as 'already marked', which decides what gets downloaded.

    This is the other half of the freeze: the marks are carried forward, so the
    fetch must recognise them and skip only those. Getting it wrong in the
    lenient direction costs a re-download; getting it wrong in the strict
    direction loses a session's data entirely.
    """

    def _existing(self, headers, pod_cell=""):
        import live_data
        b = book(tab([("a@x.com", "919000000001", pod_cell,
                       ["Present"] * len(headers))], headers))
        return live_data._existing_sessions(b)[("CAP", 35)]

    def _queued(self, existing, folder):
        """The decision fetch_new_attendees makes for one session folder."""
        cand = ac.folder_keys(folder)
        return not (cand and (cand & existing))          # True = would download

    def test_one_marked_pod_does_not_suppress_the_other_ten(self):
        """The bug the POD in the key exists to prevent.

        23 Aug 2026 ran eleven POD sessions for B35. Keyed on the date alone,
        marking Techies made the fetch skip Finance, Data, Educators and the
        rest - a whole day of sessions missing from a green run.
        """
        ex = self._existing(["2026_08_23 | Techies"], pod_cell="Techies")
        self.assertFalse(self._queued(
            ex, "2026-08-23 - AI CAP B35 - Techies - Python with AI"))
        for pod in ("Finance", "Data", "Educators", "Healthcare"):
            self.assertTrue(
                self._queued(ex, f"2026-08-23 - AI CAP B35 - {pod} - A topic"),
                f"{pod} on an already-marked date must still be downloaded")

    def test_a_whole_batch_session_on_a_marked_pod_date_is_still_queued(self):
        ex = self._existing(["2026_08_23 | Techies"], pod_cell="Techies")
        self.assertTrue(self._queued(
            ex, "2026-08-23 - AI CAP B35 - Office Productivity"))

    def test_an_unspaced_pod_in_a_folder_name_is_still_recognised(self):
        """Measured: 53 of 609 already-marked sessions were re-downloaded and
        re-marked on every run because the folder welds the POD to the batch
        with no spaces - so those columns were never frozen at all."""
        ex = self._existing(["2026_08_30 | Techies"], pod_cell="Techies")
        self.assertFalse(self._queued(
            ex, "2026-08-30 - AI CAP B37 8PM-Techis - AI-Powered Testing"))
        ex2 = self._existing(["2026_08_30 | Generalist"], pod_cell="Generalist")
        self.assertFalse(self._queued(ex2, "2026-08-30 - AICAPB35, B36-Generalist"))

    def test_a_session_a_year_later_is_not_suppressed_by_last_years_column(self):
        ex = self._existing(["2026_05_09"])
        self.assertTrue(self._queued(ex, "2027-05-09 - AI CAP B35 - A topic"))

    def test_a_legacy_header_with_no_year_still_suppresses_its_own_folder(self):
        # B17-B28 columns inside the roster Sheet read '9th May', not a full
        # date. The year-blind form has to keep working or those 133 frozen
        # columns would be re-downloaded every run forever.
        ex = self._existing(["9th May"])
        self.assertIn(("05_09", ""), ex)
        self.assertFalse(self._queued(ex, "2026-05-09 - AI CAP B35 - A topic"))

    def test_a_marker_written_header_keys_on_the_full_date_only(self):
        """One form per column, not both.

        Adding the year-blind key as well made '9th May' (2025) and
        '2026_05_09' the same key: with setdefault one of them was dropped from
        the carry AND still suppressed its own re-download, so on a freeze the
        session was gone for good.
        """
        self.assertEqual(self._existing(["2026_08_23"]), {("2026_08_23", "")})
        self.assertEqual(self._existing(["9th May"]), {("05_09", "")})

    def test_two_years_on_the_same_day_are_two_sessions(self):
        ex = self._existing(["9th May", "2026_05_09"])
        self.assertEqual(ex, {("05_09", ""), ("2026_05_09", "")})


class TestReport(unittest.TestCase):
    def test_the_summary_says_what_happened(self):
        prev = book(tab([("a@x.com", "919000000001", "", ["Present"])],
                        ["2026_08_23"]))
        fresh = book(tab([("a@x.com", "919000000001", "", [])]))
        _out, rep = cf.merge_marks(fresh, prev)
        line = cf.summary(rep)
        self.assertIn("1 column(s) carried", line)
        self.assertIn("1 student(s) matched", line)

    def test_carrying_nothing_says_so_rather_than_looking_healthy(self):
        self.assertIn("carried nothing", cf.summary({"carried": 0, "warnings": []}))


def report_csv(rows, wid="9912345678", date="09/13/2026"):
    """A real Zoom Attendee Report, same shape as tests/test_attendee_format."""
    out = ["Attendee Report",
           "Report generated time,09/13/2026 05:29:06 AM",
           "",
           "Topic,Webinar ID,Actual Start Time,Unique Viewers",
           f"Some Session,{wid},{date} 19:00,{len(rows)}",
           "",
           "Attendee Details",
           "Attended,First Name,Last Name,Email,Phone,Join Time,Leave Time,"
           "Time in Session (minutes)"]
    for email, phone in rows:
        out.append(f"Yes,First,Last,{email},{phone},{date} 19:02:00,"
                   f"{date} 20:31:00,89")
    return "\n".join(out).encode()


class TestFidelity(unittest.TestCase):
    """The failures that would corrupt history quietly rather than loudly."""

    def test_a_renamed_tab_still_carries_its_history(self):
        # The marker matches batches through _sheet_key precisely because tab
        # names drift. Matching on the NAME would drop a whole batch in silence.
        prev = book(tab([("a@x.com", "919000000001", "", ["Present"])],
                        ["2026_08_23"], sheet="AI CAP B37"))
        fresh = book(tab([("a@x.com", "919000000001", "", [])], [],
                         sheet="AI CAP B37 8PM"))
        out, rep = cf.merge_marks(fresh, prev)
        _hdr, rows = read(out, "AI CAP B37 8PM")
        self.assertEqual(rows["a@x.com"]["2026_08_23"], "Present")
        self.assertEqual(rep["carried"], 1)

    def test_a_batch_tab_that_vanished_from_the_roster_is_reported(self):
        # Four different people own the source Sheets (CLAUDE section 4d). A
        # deleted tab takes its whole history with it, so it must not be silent.
        prev = book(tab([("a@x.com", "919000000001", "", ["Present"])],
                        ["2026_08_23"]),
                    tab([("g@x.com", "919000000002", "", ["Present"])],
                        ["2026_08_23"], sheet="AI CAP B34"))
        fresh = book(tab([("a@x.com", "919000000001", "", [])]))
        _out, rep = cf.merge_marks(fresh, prev)
        self.assertTrue(any("B34" in w for w in rep["warnings"]),
                        rep["warnings"])

    def test_students_whose_history_is_dropped_are_counted_per_tab(self):
        # The one counter that reveals "we lost N students' marks". Comparing a
        # per-tab count against a running total made it read 0 from the second
        # tab on.
        prev = book(tab([("a@x.com", "919000000001", "", ["Present"]),
                         ("gone1@x.com", "919000000011", "", ["Present"])],
                        ["2026_08_23"]),
                    tab([("b@x.com", "919000000002", "", ["Present"]),
                         ("gone2@x.com", "919000000012", "", ["Present"])],
                        ["2026_08_24"], sheet="AI CAP B36"))
        fresh = book(tab([("a@x.com", "919000000001", "", [])]),
                     tab([("b@x.com", "919000000002", "", [])], [],
                         sheet="AI CAP B36"))
        _out, rep = cf.merge_marks(fresh, prev)
        self.assertEqual(rep["unmatched_prev"], 2)
        self.assertEqual(rep["matched"], 2)

    def test_an_email_change_is_not_reported_as_a_lost_student(self):
        # The case the design exists to handle must not also raise the warning
        # that says it failed. The old email key is never touched, so counting
        # KEYS rather than students reported this student as not carried.
        prev = book(tab([("old@x.com", "919886440098", "", ["Present"])],
                        ["2026_08_23"]))
        fresh = book(tab([("new@x.com", "919886440098", "", [])]))
        _out, rep = cf.merge_marks(fresh, prev)
        self.assertEqual(rep["matched"], 1)
        self.assertEqual(rep["unmatched_prev"], 0)
        self.assertEqual(rep["warnings"], [])

    def test_a_short_phone_number_keeps_its_marks(self):
        """_phone_hit checks the WHOLE number before the last ten digits,
        "kept for numbers too short to have a last-10 form". Indexing only the
        last ten made such a student look new - and a new student is written
        'Absent', so their Presents were overwritten for good."""
        # The EMAIL must differ, or it matches first and the whole-number phone
        # branch this test exists to pin is never exercised.
        prev = book(tab([("old@x.com", "98765432", "", ["Present"])],
                        ["2026_08_23"]))
        fresh = book(tab([("new@x.com", "98765432", "", [])]))
        out, rep = cf.merge_marks(fresh, prev)
        _hdr, rows = read(out)
        self.assertEqual(rows["new@x.com"]["2026_08_23"], "Present")
        self.assertEqual(rep["new_students"], 0)
        self.assertEqual(rep["matched"], 1)

    def test_a_duplicated_student_keeps_the_present(self):
        # Two rows for one person, one Present one Absent: the marker's rule is
        # an OR over identity, so Present is the true answer - and it must not
        # depend on which row came first.
        for order in ((("dup@x.com", "919000000001", "", ["Present"]),
                       ("dup@x.com", "919000000001", "", ["Absent"])),
                      (("dup@x.com", "919000000001", "", ["Absent"]),
                       ("dup@x.com", "919000000001", "", ["Present"]))):
            prev = book(tab(list(order), ["2026_08_23"]))
            fresh = book(tab([("dup@x.com", "919000000001", "", [])]))
            _hdr, rows = read(cf.merge_marks(fresh, prev)[0])
            self.assertEqual(rows["dup@x.com"]["2026_08_23"], "Present")

    def test_carried_columns_keep_the_previous_order_across_a_year_boundary(self):
        # Sorting by (month, day) puts December after January and scrambles
        # SessionIdx and the Roster tab's columns.
        prev = book(tab([("a@x.com", "919000000001", "", ["P1", "P2", "P3"])],
                        ["2025_12_20", "2026_01_10", "2026_02_14"]))
        fresh = book(tab([("a@x.com", "919000000001", "", [])]))
        hdr, _rows = read(cf.merge_marks(fresh, prev)[0])
        got = [h for h in hdr if h and ac.session_key(h)]
        self.assertEqual(got, ["2025_12_20", "2026_01_10", "2026_02_14"])

    def test_the_topic_row_is_carried_for_a_two_row_header(self):
        wb = Workbook()
        wb.remove(wb.active)
        ws = wb.create_sheet("AI CAP B35")
        ws.append([None] * 11 + ["2026_08_23 | Techies"])      # dates row
        ws.append(list(FIXED) + ["Python with AI"])            # header + topic
        ws.append([91, "919000000001", "a@x.com", 91, "9", "a@x.com",
                   "AI CAP B35", 1, "Full Paid", "BDA", "Techies", "Present"])
        buf = io.BytesIO()
        wb.save(buf)
        fresh = Workbook()
        fresh.remove(fresh.active)
        fs = fresh.create_sheet("AI CAP B35")
        fs.append([None] * 11)
        fs.append(list(FIXED))
        fs.append([91, "919000000001", "a@x.com", 91, "9", "a@x.com",
                   "AI CAP B35", 1, "Full Paid", "BDA", "Techies"])
        fbuf = io.BytesIO()
        fresh.save(fbuf)
        out, rep = cf.merge_marks(fbuf.getvalue(), buf.getvalue())
        self.assertEqual(rep["carried"], 1)
        got = load_workbook(io.BytesIO(out), data_only=True)["AI CAP B35"]
        col = next(c for c in range(1, got.max_column + 1)
                   if ac.session_key(got.cell(1, c).value))
        self.assertEqual(got.cell(1, col).value, "2026_08_23 | Techies")
        self.assertEqual(got.cell(2, col).value, "Python with AI")
        self.assertEqual(got.cell(3, col).value, "Present")

    def test_merging_twice_changes_nothing(self):
        prev = book(tab([("a@x.com", "919000000001", "", ["Present"])],
                        ["2026_08_23"]))
        fresh = book(tab([("a@x.com", "919000000001", "", []),
                          ("new@x.com", "919000000009", "", [])]))
        once, r1 = cf.merge_marks(fresh, prev)
        twice, r2 = cf.merge_marks(once, prev)
        self.assertEqual(r2["carried"], 0)
        self.assertEqual(r2["already_present"], 1)
        self.assertEqual(read(once)[1], read(twice)[1])


class TestWithTheMarker(unittest.TestCase):
    """The merge and the marker have to agree about what a column IS."""

    def test_one_new_file_adds_exactly_one_column_and_touches_nothing_else(self):
        prev = book(tab([("a@x.com", "919000000001", "", ["Present"]),
                         ("b@x.com", "919000000002", "", ["Absent"])],
                        ["2026_08_23"]))
        fresh = book(tab([("a@x.com", "919000000001", "", []),
                          ("b@x.com", "919000000002", "", [])]))
        base, _rep = cf.merge_marks(fresh, prev)
        marked, report, warnings = ac.process_files(
            base, None,
            [("2026-09-13 - AI CAP B35 - A topic/attendee_9912345678_2026_09_13.csv",
              report_csv([("a@x.com", "919000000001")]))],
            values_only=True)
        self.assertEqual([r["kind"] for r in report], ["NEW"])
        # the only warning is the expected "no L2 in this fixture" note
        self.assertEqual([w for w in warnings if ac.ZERO_ATTENDEE_TAG in w], [])
        hdr, rows = read(marked)
        self.assertEqual([h for h in hdr if h and ac.session_key(h)],
                         ["2026_08_23", "2026_09_13"])
        # last week frozen, this week marked
        self.assertEqual(rows["a@x.com"]["2026_08_23"], "Present")
        self.assertEqual(rows["b@x.com"]["2026_08_23"], "Absent")
        self.assertEqual(rows["a@x.com"]["2026_09_13"], "Present")
        self.assertEqual(rows["b@x.com"]["2026_09_13"], "Absent")

    def test_a_second_file_for_a_carried_session_re_marks_rather_than_duplicating(self):
        prev = book(tab([("a@x.com", "919000000001", "", ["Absent"])],
                        ["2026_09_13"]))
        fresh = book(tab([("a@x.com", "919000000001", "", [])]))
        base, _rep = cf.merge_marks(fresh, prev)
        marked, report, _w = ac.process_files(
            base, None,
            [("2026-09-13 - AI CAP B35 - A topic/attendee_9912345678_2026_09_13.csv",
              report_csv([("a@x.com", "919000000001")]))],
            values_only=True)
        self.assertEqual([r["kind"] for r in report], ["re-mark"])
        hdr, rows = read(marked)
        self.assertEqual([h for h in hdr if h and ac.session_key(h)], ["2026_09_13"])
        self.assertEqual(rows["a@x.com"]["2026_09_13"], "Present")

    def test_a_frozen_column_is_left_alone_even_when_its_files_arrive(self):
        """The bug the `frozen` set exists to prevent, measured on real data.

        One session can sit in TWO folders whose names disagree — the POD-named
        one is skipped as already marked, the other is downloaded, and the marker
        used to rewrite the carried column from that second copy of the same
        Zoom export. The copies differ by a row or two: B35's 23 Aug Educators
        session went 63 -> 60 present and Sales/Marketing/HR 88 -> 87, and
        GATE 6 refused to publish. Freezing at the point of WRITE makes it
        irrelevant which copy the fetch hands over.
        """
        prev = book(tab([("a@x.com", "919000000001", "", ["Present"])],
                        ["2026_09_13"]))
        fresh = book(tab([("a@x.com", "919000000001", "", [])]))
        base, rep = cf.merge_marks(fresh, prev)
        frozen = rep["_carried_keys"]
        self.assertEqual(frozen, {("CAP", 35): {("2026_09_13", "")}})
        # a copy of that very session turns up, with FEWER attendees
        marked, report, _w = ac.process_files(
            base, None,
            [("2026-09-13 - AI CAP B35 - A topic/attendee_9912345678_2026_09_13.csv",
              report_csv([("someone.else@x.com", "919000009999")]))],
            values_only=True, frozen=frozen)
        self.assertEqual([r["kind"] for r in report], ["frozen"])
        _hdr, rows = read(marked)
        self.assertEqual(rows["a@x.com"]["2026_09_13"], "Present")   # untouched

    def test_a_frozen_batch_does_not_freeze_a_different_batch(self):
        prev = book(tab([("a@x.com", "919000000001", "", ["Present"])],
                        ["2026_09_13"]))
        fresh = book(tab([("a@x.com", "919000000001", "", [])]),
                     tab([("n@x.com", "919000000009", "", [])], [],
                         sheet="AI CAP B40"))
        base, rep = cf.merge_marks(fresh, prev)
        # the same date for B40, which carried nothing, must still be marked
        marked, report, _w = ac.process_files(
            base, None,
            [("2026-09-13 - AI CAP B40 - A topic/attendee_8812345678_2026_09_13.csv",
              report_csv([("n@x.com", "919000000009")]))],
            values_only=True, frozen=rep["_carried_keys"])
        self.assertEqual([r["kind"] for r in report], ["NEW"])
        _hdr, rows40 = read(marked, "AI CAP B40")
        self.assertEqual(rows40["n@x.com"]["2026_09_13"], "Present")

    def test_a_session_a_year_later_does_not_overwrite_the_carried_column(self):
        """A real 2027 session, with a real attendee, must not land on the 2026
        column of the same day.

        The first version of this test used an EMPTY 2027 report, so
        process_files skipped the session entirely and the assertion passed
        whether or not the year rule worked at all. The attendee below is what
        makes it a test.
        """
        prev = book(tab([("a@x.com", "919000000001", "", ["Present"])],
                        ["2026_09_13"]))
        fresh = book(tab([("a@x.com", "919000000001", "", [])]))
        base, _rep = cf.merge_marks(fresh, prev)
        marked, report, _w = ac.process_files(
            base, None,
            [("2027-09-13 - AI CAP B35 - A topic/attendee_9912345678_2027_09_13.csv",
              report_csv([("a@x.com", "919000000001")], date="09/13/2027"))],
            values_only=True)
        self.assertEqual([r["kind"] for r in report], ["NEW"])   # a real mark
        hdr, rows = read(marked)
        self.assertEqual([h for h in hdr if h and ac.session_key(h)],
                         ["2026_09_13", "2027_09_13"])
        self.assertEqual(rows["a@x.com"]["2026_09_13"], "Present")
        self.assertEqual(rows["a@x.com"]["2027_09_13"], "Present")


if __name__ == "__main__":
    unittest.main()
